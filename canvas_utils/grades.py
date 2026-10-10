"""Assignment grade export logic for Canvas assignments."""

from __future__ import annotations

import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from canvas_utils.api_client import CanvasAPIError, CanvasClient
from canvas_utils.rubrics import (
    criterion_id_aliases as _criterion_id_aliases,
    extract_rubric_criteria as _extract_rubric_criteria,
    find_rubric_rating as _find_rubric_rating,
    identifier_or_none as _identifier_or_none,
    index_rubric_criteria as _index_rubric_criteria,
    rubric_criterion_description as _rubric_criterion_description,
)

GRADE_HEADERS = [
    "Assignment_ID",
    "Assignment_Name",
    "Student_ID",
    "Student_Name",
    "Student_Login_ID",
    "Score",
    "Grade",
    "Entered_Grade",
    "Workflow_State",
    "Grader_ID",
    "Grader_Name",
    "Grader_Login_ID",
    "Grader_SIS_ID",
    "Grader_Type",
]

MARKER_STATS_HEADERS = [
    "Grader_ID",
    "Grader_Name",
    "Grader_Login_ID",
    "Grader_SIS_ID",
    "Grader_Type",
    "Submission_Count",
    "Scored_Submission_Count",
    "Unscored_Submission_Count",
    "Average_Score",
    "Minimum_Score",
    "Maximum_Score",
]

RUBRIC_ASSESSMENT_HEADERS = [
    "Assignment_ID",
    "Assignment_Name",
    "Student_ID",
    "Student_Name",
    "Student_Login_ID",
    "Grader_ID",
    "Grader_Name",
    "Grader_Login_ID",
    "Grader_SIS_ID",
    "Grader_Type",
    "Criterion_ID",
    "Criterion_Description",
    "Rating_ID",
    "Selected_Descriptor",
    "Points",
    "Comments",
]

AUTO_GRADED_LABEL = "Auto-graded"
UNGRADED_LABEL = "Ungraded"
UNRESOLVED_GRADER_LABEL = "Unresolvable grader"


def export_assignment_grades(
    client: CanvasClient,
    course_id: int,
    assignment_id: int | None,
    output: str | Path = ".",
    *,
    include_rubric_assessment: bool = False,
) -> list[Path]:
    """Export assignment grades with grader identity to XLSX."""
    target_assignments = _resolve_assignment_records(
        client,
        course_id,
        assignment_id,
        include_rubric=include_rubric_assessment,
    )
    output_path = Path(output)
    grader_cache: dict[int, dict[str, Any]] = {}
    written_files: list[Path] = []

    for current_assignment_id, assignment_name, rubric_criteria in target_assignments:
        if include_rubric_assessment:
            submissions = client.list_submissions(
                course_id=course_id,
                assignment_id=current_assignment_id,
                include_user=True,
                include_rubric_assessment=True,
            )
        else:
            submissions = client.list_submissions(
                course_id=course_id,
                assignment_id=current_assignment_id,
                include_user=True,
            )

        rows = [
            _map_grade_row(
                submission=submission,
                course_id=course_id,
                fallback_assignment_id=current_assignment_id,
                assignment_name=assignment_name,
                client=client,
                grader_cache=grader_cache,
            )
            for submission in submissions
        ]
        rubric_rows = (
            [
                rubric_row
                for submission, grade_row in zip(submissions, rows)
                for rubric_row in _map_rubric_assessment_rows(
                    submission=submission,
                    grade_row=grade_row,
                    rubric_criteria=rubric_criteria,
                )
            ]
            if include_rubric_assessment
            else None
        )
        xlsx_path = _resolve_xlsx_output_path(
            output_path=output_path,
            assignment_id=assignment_id,
            current_assignment_id=current_assignment_id,
            assignment_name=assignment_name,
        )
        _write_workbook(xlsx_path, rows, rubric_rows)
        written_files.append(xlsx_path)

    return written_files


def _map_grade_row(
    submission: dict[str, Any],
    course_id: int,
    fallback_assignment_id: int,
    assignment_name: str,
    client: CanvasClient,
    grader_cache: dict[int, dict[str, Any]],
) -> dict[str, Any]:
    grader_id = _coerce_int(submission.get("grader_id"))
    grader_type, grader_name, grader_login_id, grader_sis_id = _resolve_grader_identity(
        course_id=course_id,
        grader_id=grader_id,
        client=client,
        grader_cache=grader_cache,
    )
    student = _submission_user(submission)

    return {
        "Assignment_ID": _coerce_positive_int(submission.get("assignment_id")) or fallback_assignment_id,
        "Assignment_Name": assignment_name,
        "Student_ID": _coerce_positive_int(submission.get("user_id")),
        "Student_Name": student.get("name"),
        "Student_Login_ID": student.get("login_id"),
        "Score": _coerce_score(submission.get("score")),
        "Grade": submission.get("grade"),
        "Entered_Grade": submission.get("entered_grade"),
        "Workflow_State": submission.get("workflow_state"),
        "Grader_ID": grader_id,
        "Grader_Name": grader_name,
        "Grader_Login_ID": grader_login_id,
        "Grader_SIS_ID": grader_sis_id,
        "Grader_Type": grader_type,
    }


def _map_rubric_assessment_rows(
    *,
    submission: dict[str, Any],
    grade_row: dict[str, Any],
    rubric_criteria: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    rubric_assessment = submission.get("rubric_assessment")
    if not isinstance(rubric_assessment, dict):
        return []

    criteria_by_id = _index_rubric_criteria(rubric_criteria)
    rows: list[dict[str, Any]] = []

    for raw_criterion_id, selection in rubric_assessment.items():
        if not isinstance(selection, dict):
            continue

        criterion_id = _identifier_or_none(raw_criterion_id)
        if criterion_id is None:
            continue

        criterion = criteria_by_id.get(criterion_id)
        rating_id = _identifier_or_none(selection.get("rating_id"))
        rating = _find_rubric_rating(criterion, rating_id)
        resolved_criterion_id = (
            _identifier_or_none(criterion.get("id")) if criterion is not None else criterion_id
        )

        rows.append(
            {
                "Assignment_ID": grade_row.get("Assignment_ID"),
                "Assignment_Name": grade_row.get("Assignment_Name"),
                "Student_ID": grade_row.get("Student_ID"),
                "Student_Name": grade_row.get("Student_Name"),
                "Student_Login_ID": grade_row.get("Student_Login_ID"),
                "Grader_ID": grade_row.get("Grader_ID"),
                "Grader_Name": grade_row.get("Grader_Name"),
                "Grader_Login_ID": grade_row.get("Grader_Login_ID"),
                "Grader_SIS_ID": grade_row.get("Grader_SIS_ID"),
                "Grader_Type": grade_row.get("Grader_Type"),
                "Criterion_ID": resolved_criterion_id,
                "Criterion_Description": _rubric_criterion_description(criterion),
                "Rating_ID": rating_id,
                "Selected_Descriptor": _selected_rating_descriptor(rating, selection),
                "Points": _coerce_score(selection.get("points")),
                "Comments": selection.get("comments"),
            }
        )

    return rows


def _selected_rating_descriptor(
    rating: dict[str, Any] | None,
    selection: dict[str, Any],
) -> str | None:
    if rating is not None:
        descriptor = _string_or_none(rating.get("description")) or _string_or_none(rating.get("long_description"))
        if descriptor is not None:
            return descriptor
    return _string_or_none(selection.get("description")) or _string_or_none(selection.get("long_description"))


def build_marker_statistics(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Build per-marker identity rows for the statistics worksheet."""
    grouped_rows: dict[tuple[Any, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        key = (
            row.get("Grader_ID"),
            row.get("Grader_Name"),
            row.get("Grader_Login_ID"),
            row.get("Grader_SIS_ID"),
            row.get("Grader_Type"),
        )
        grouped_rows[key].append(row)

    statistics_rows: list[dict[str, Any]] = []
    for key in sorted(grouped_rows, key=_marker_sort_key):
        statistics_rows.append(
            {
                "Grader_ID": key[0],
                "Grader_Name": key[1],
                "Grader_Login_ID": key[2],
                "Grader_SIS_ID": key[3],
                "Grader_Type": key[4],
            }
        )

    return statistics_rows


def _resolve_grader_identity(
    *,
    course_id: int,
    grader_id: int | None,
    client: CanvasClient,
    grader_cache: dict[int, dict[str, Any]],
) -> tuple[str, str | None, str | None, str | None]:
    if grader_id is None:
        return "ungraded", UNGRADED_LABEL, None, None

    if grader_id < 0:
        return "automatic", AUTO_GRADED_LABEL, None, None

    grader = grader_cache.get(grader_id)
    if grader is None:
        try:
            grader = client.get_user(grader_id)
        except CanvasAPIError as exc:
            if exc.status_code != 404:
                raise
            try:
                grader = client.get_course_user(course_id, grader_id)
            except CanvasAPIError as course_exc:
                if course_exc.status_code == 404:
                    return "unresolvable", UNRESOLVED_GRADER_LABEL, None, None
                raise
        grader_cache[grader_id] = grader

    return (
        "human",
        _string_or_none(grader.get("name")),
        _string_or_none(grader.get("login_id")),
        _string_or_none(grader.get("sis_user_id")),
    )


def _resolve_assignment_records(
    client: CanvasClient,
    course_id: int,
    assignment_id: int | None,
    *,
    include_rubric: bool = False,
) -> list[tuple[int, str, list[dict[str, Any]]]]:
    if assignment_id is not None:
        assignments = client.list_assignments(course_id)
        matched_name: str | None = None
        for assignment in assignments:
            current_id = _coerce_positive_int(assignment.get("id"))
            if current_id == assignment_id:
                matched_name = _assignment_name(assignment)
                break

        assignment_name = matched_name or f"assignment-{assignment_id}"
        if include_rubric:
            rubric_assignment = client.get_assignment(
                course_id=course_id,
                assignment_id=assignment_id,
                include_rubric=True,
            )
            name = _assignment_name(rubric_assignment)
            if name == "assignment" and assignment_name != "assignment":
                name = assignment_name
            return [(assignment_id, name, _extract_rubric_criteria(rubric_assignment))]

        return [(assignment_id, assignment_name, [])]

    assignments = client.list_assignments(course_id, include_rubric=include_rubric)
    return [
        (
            current_id,
            _assignment_name(assignment),
            _extract_rubric_criteria(assignment) if include_rubric else [],
        )
        for assignment in assignments
        if (current_id := _coerce_positive_int(assignment.get("id"))) is not None
    ]


def _resolve_xlsx_output_path(
    output_path: Path,
    assignment_id: int | None,
    current_assignment_id: int,
    assignment_name: str,
) -> Path:
    export_date = datetime.now().strftime("%Y-%m-%d")
    file_name = f"canvas_grades_{current_assignment_id}_{_slugify_filename_part(assignment_name)}_{export_date}.xlsx"

    if assignment_id is None:
        if output_path.suffix.lower() == ".xlsx":
            raise ValueError("When --assignment is 'all', --output must be a directory, not an XLSX file path.")
        output_path.mkdir(parents=True, exist_ok=True)
        return output_path / file_name

    if output_path.suffix.lower() == ".xlsx":
        output_path.parent.mkdir(parents=True, exist_ok=True)
        return output_path

    output_path.mkdir(parents=True, exist_ok=True)
    return output_path / file_name


def _write_workbook(
    xlsx_path: Path,
    rows: list[dict[str, Any]],
    rubric_rows: list[dict[str, Any]] | None = None,
) -> None:
    workbook = Workbook()
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    grades_sheet = workbook.active
    grades_sheet.title = "Grades"
    _append_sheet_rows(grades_sheet, GRADE_HEADERS, rows)

    stats_sheet = workbook.create_sheet(title="Marker Statistics")
    _append_marker_statistics_rows(stats_sheet, build_marker_statistics(rows))

    if rubric_rows is not None:
        rubric_sheet = workbook.create_sheet(title="Rubric Assessments")
        _append_sheet_rows(rubric_sheet, RUBRIC_ASSESSMENT_HEADERS, rubric_rows)

    workbook.save(xlsx_path)


def _sanitize_cell_value(val: Any) -> Any:
    if isinstance(val, str) and val.startswith("="):
        return f"'{val}"
    return val


def _append_sheet_rows(worksheet: Any, headers: list[str], rows: list[dict[str, Any]]) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append([_sanitize_cell_value(row.get(header)) for header in headers])


def _append_marker_statistics_rows(worksheet: Any, rows: list[dict[str, Any]]) -> None:
    worksheet.append(MARKER_STATS_HEADERS)
    for row_number, row in enumerate(rows, start=2):
        worksheet.append(
            [
                _sanitize_cell_value(row.get("Grader_ID")),
                _sanitize_cell_value(row.get("Grader_Name")),
                _sanitize_cell_value(row.get("Grader_Login_ID")),
                _sanitize_cell_value(row.get("Grader_SIS_ID")),
                _sanitize_cell_value(row.get("Grader_Type")),
                _count_submissions_formula(row_number),
                _count_scored_submissions_formula(row_number),
                f"=F{row_number}-G{row_number}",
                _average_score_formula(row_number),
                _minimum_score_formula(row_number),
                _maximum_score_formula(row_number),
            ]
        )


def _submission_user(submission: dict[str, Any]) -> dict[str, Any]:
    user = submission.get("user")
    return user if isinstance(user, dict) else {}


def _assignment_name(assignment: dict[str, Any]) -> str:
    name = assignment.get("name")
    if isinstance(name, str) and name.strip():
        return name.strip()
    return "assignment"


def _slugify_filename_part(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_")
    return slug.lower() or "assignment"


def _coerce_positive_int(value: object) -> int | None:
    if isinstance(value, int) and value > 0:
        return value
    if isinstance(value, str) and value.isdigit():
        parsed = int(value)
        return parsed if parsed > 0 else None
    return None


def _coerce_int(value: object) -> int | None:
    if isinstance(value, int):
        return value
    if isinstance(value, str) and re.fullmatch(r"-?\d+", value):
        return int(value)
    return None


def _coerce_score(value: object) -> float | int | None:
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _count_submissions_formula(row_number: int) -> str:
    criteria = _marker_criteria_formula(row_number)
    return f"=COUNTIFS({criteria})"


def _count_scored_submissions_formula(row_number: int) -> str:
    criteria_terms = _marker_criteria_terms(row_number)
    score_column = _grade_range("Score")
    return f'=SUMPRODUCT(({criteria_terms})*ISNUMBER({score_column}))'


def _average_score_formula(row_number: int) -> str:
    criteria = _marker_criteria_formula(row_number)
    score_column = _grade_range("Score")
    return f'=IF(G{row_number}=0,"",ROUND(AVERAGEIFS({score_column},{criteria}),2))'


def _minimum_score_formula(row_number: int) -> str:
    criteria = _marker_criteria_formula(row_number)
    score_column = _grade_range("Score")
    return f'=IF(G{row_number}=0,"",MINIFS({score_column},{criteria}))'


def _maximum_score_formula(row_number: int) -> str:
    criteria = _marker_criteria_formula(row_number)
    score_column = _grade_range("Score")
    return f'=IF(G{row_number}=0,"",MAXIFS({score_column},{criteria}))'


def _marker_criteria_formula(row_number: int) -> str:
    return ",".join(
        [
            f"{_grade_range('Grader_ID')},A{row_number}",
            f"{_grade_range('Grader_Name')},B{row_number}",
            f"{_grade_range('Grader_Login_ID')},C{row_number}",
            f"{_grade_range('Grader_SIS_ID')},D{row_number}",
            f"{_grade_range('Grader_Type')},E{row_number}",
        ]
    )


def _marker_criteria_terms(row_number: int) -> str:
    return "*".join(
        [
            f"({_grade_range('Grader_ID')}=A{row_number})",
            f"({_grade_range('Grader_Name')}=B{row_number})",
            f"({_grade_range('Grader_Login_ID')}=C{row_number})",
            f"({_grade_range('Grader_SIS_ID')}=D{row_number})",
            f"({_grade_range('Grader_Type')}=E{row_number})",
        ]
    )


def _grade_range(header: str) -> str:
    column_letter = get_column_letter(GRADE_HEADERS.index(header) + 1)
    return f"Grades!${column_letter}$2:${column_letter}$1048576"


def _string_or_none(value: object) -> str | None:
    if isinstance(value, str) and value.strip():
        return value
    return None


def _marker_sort_key(marker_key: tuple[Any, ...]) -> tuple[str, str]:
    grader_type = str(marker_key[4] or "")
    grader_name = str(marker_key[1] or "")
    return grader_type, grader_name.casefold()
