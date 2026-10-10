"""Controlled XLSX gradebook creation, validation, and upload workflows."""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
import uuid
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from canvas_utils.api_client import CanvasAPIError, CanvasClient, CanvasTransportError
from canvas_utils.rubrics import (
    criterion_id_aliases,
    extract_rubric_criteria,
    find_rubric_rating,
    identifier_or_none,
    index_rubric_criteria,
    rubric_criterion_description,
    rubric_rating_description,
    rubric_ratings,
)


GRADEBOOK_SCHEMA_VERSION = "1"
UPLOAD_JOURNAL_SCHEMA_VERSION = "1"

INSTRUCTIONS_SHEET = "Instructions"
GRADES_SHEET = "Grades"
RUBRIC_SHEET = "Rubric"
RUBRIC_CATALOG_SHEET = "Rubric Catalog"
MAPPING_SHEET = "Canvas Mapping"
METADATA_SHEET = "Canvas Metadata"
SNAPSHOT_SHEET = "Canvas Snapshot"
RUBRIC_LISTS_SHEET = "Rubric Lists"

GRADE_HEADERS = [
    "Canvas_User_ID",
    "Student_Name",
    "Student_Login_ID",
    "Current_Grade",
    "New_Grade",
    "Grade_Action",
    "New_Submission_Comment",
    "Comment_Action",
]

RUBRIC_HEADERS = [
    "Canvas_User_ID",
    "Student_Name",
    "Criterion_ID",
    "Criterion",
    "Current_Descriptor",
    "Current_Points",
    "Current_Criterion_Comment",
    "Selected_Descriptor",
    "Selected_Rating_ID",
    "Points",
    "Criterion_Comment",
    "Action",
]

RUBRIC_CATALOG_HEADERS = [
    "Criterion_ID",
    "Criterion",
    "Rating_ID",
    "Descriptor_Selection",
    "Descriptor",
    "Default_Points",
]

MAPPING_HEADERS = ["Field", "Sheet", "Column"]
METADATA_HEADERS = ["Key", "Value"]
SNAPSHOT_HEADERS = [
    "Record_Type",
    "Canvas_User_ID",
    "Criterion_ID",
    "Submission_Attempt",
    "Submitted_At",
    "Grade",
    "Rating_ID",
    "Points",
    "Criterion_Comment",
]

GRADE_MAPPING_FIELDS = (
    "grades.student_id",
    "grades.new_grade",
    "grades.grade_action",
    "grades.submission_comment",
    "grades.comment_action",
)
RUBRIC_MAPPING_FIELDS = (
    "rubric.student_id",
    "rubric.criterion_id",
    "rubric.selected_descriptor",
    "rubric.points",
    "rubric.criterion_comment",
    "rubric.action",
)
OPTIONAL_RUBRIC_MAPPING_FIELDS = ("rubric.selected_rating_id",)
ALL_MAPPING_FIELDS = frozenset(
    (*GRADE_MAPPING_FIELDS, *RUBRIC_MAPPING_FIELDS, *OPTIONAL_RUBRIC_MAPPING_FIELDS)
)


@dataclass(frozen=True)
class _RubricSnapshot:
    rating_id: object
    points: object
    comment: object


@dataclass(frozen=True)
class _SubmissionSnapshot:
    attempt: object
    submitted_at: object


@dataclass(frozen=True)
class _ChangeBaseline:
    grade: object
    rubric_assessments: dict[str, _RubricSnapshot]
    submission: _SubmissionSnapshot


@dataclass(frozen=True)
class GradebookChange:
    """One student's validated Canvas submission update."""

    user_id: int
    student_name: str | None
    posted_grade: str | int | float | None = None
    submission_comment: str | None = None
    submission_comment_attempt: int | None = None
    rubric_assessment: dict[str, dict[str, Any]] = field(default_factory=dict)
    baseline: _ChangeBaseline | None = field(default=None, repr=False, compare=False)


@dataclass(frozen=True)
class GradebookPlan:
    """A fully validated set of gradebook changes ready for dry run or upload."""

    course_id: int
    assignment_id: int
    template_id: str
    total_students: int
    unchanged_students: int
    changes: tuple[GradebookChange, ...]


@dataclass(frozen=True)
class GradebookUploadResult:
    """Outcome counts for a gradebook dry run or upload."""

    planned: int
    applied: int
    failed: int
    conflicted: int
    indeterminate: int
    journal_skipped: int
    journal_errors: int
    skipped: int
    failures: tuple[str, ...]
    outcomes: tuple["GradebookUploadOutcome", ...]


@dataclass(frozen=True)
class GradebookUploadOutcome:
    """One execution outcome that can be recorded in an upload journal."""

    change_id: str
    user_id: int
    status: str
    has_submission_comment: bool


class GradebookValidationError(ValueError):
    """Raised when a gradebook cannot be safely uploaded."""

    def __init__(self, errors: list[str]) -> None:
        self.errors = tuple(errors)
        message = "Gradebook validation failed:\n" + "\n".join(f"- {error}" for error in errors)
        super().__init__(message)


class GradebookUploadLockError(ValueError):
    """Raised when another local process is already uploading this gradebook."""


class GradebookUploadLock:
    """An advisory, process-scoped lock held for an executed gradebook upload."""

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self._handle: Any | None = None

    def __enter__(self) -> "GradebookUploadLock":
        self.path.parent.mkdir(parents=True, exist_ok=True)
        handle = self.path.open("a+", encoding="utf-8")
        try:
            handle.seek(0, 2)
            if handle.tell() == 0:
                handle.write("0")
                handle.flush()
            handle.seek(0)
            if os.name == "nt":
                import msvcrt

                msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl

                fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError as exc:
            handle.close()
            raise GradebookUploadLockError(
                f"Another gradebook upload is already using lock '{self.path}'."
            ) from exc
        self._handle = handle
        return self

    def __exit__(self, exc_type: object, exc_value: object, traceback: object) -> None:
        if self._handle is not None:
            self._handle.close()
            self._handle = None


class GradebookUploadJournal:
    """Persist pending and completed gradebook requests without feedback text."""

    _RECORD_STATUSES = frozenset({"pending", "applied", "indeterminate"})

    def __init__(self, path: Path, records: dict[str, dict[str, object]]) -> None:
        self.path = path
        self._records = records

    @classmethod
    def load(cls, path: str | Path) -> "GradebookUploadJournal":
        """Load a journal, or create an in-memory journal for a new path."""
        journal_path = Path(path)
        if not journal_path.exists():
            return cls(journal_path, {})
        try:
            payload = json.loads(journal_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ValueError(f"Could not read gradebook upload journal '{journal_path}': {exc}") from exc
        if not isinstance(payload, dict):
            raise ValueError(f"Gradebook upload journal '{journal_path}' must contain a JSON object.")
        if payload.get("schema_version") != UPLOAD_JOURNAL_SCHEMA_VERSION:
            raise ValueError(f"Gradebook upload journal '{journal_path}' has an unsupported schema version.")
        raw_records = payload.get("records")
        if not isinstance(raw_records, dict):
            raise ValueError(f"Gradebook upload journal '{journal_path}' has invalid records.")

        records: dict[str, dict[str, object]] = {}
        for change_id, record in raw_records.items():
            if not isinstance(change_id, str) or not re.fullmatch(r"[0-9a-f]{64}", change_id):
                raise ValueError(f"Gradebook upload journal '{journal_path}' has an invalid change ID.")
            if not isinstance(record, dict) or record.get("status") not in cls._RECORD_STATUSES:
                raise ValueError(f"Gradebook upload journal '{journal_path}' has an invalid record.")
            records[change_id] = record
        return cls(journal_path, records)

    def excluded_change_ids(self, *, retry_indeterminate: bool = False) -> set[str]:
        """Return changes that should not be resent by default."""
        return {
            change_id
            for change_id, record in self._records.items()
            if record.get("status") == "applied"
            or (
                record.get("status") in {"pending", "indeterminate"}
                and not retry_indeterminate
            )
        }

    def prepare(self, outcome: GradebookUploadOutcome) -> None:
        """Durably mark a request pending before it is sent to Canvas."""
        self._set_record(outcome, status="pending")

    def complete(self, outcome: GradebookUploadOutcome) -> None:
        """Replace a pending record with an applied or indeterminate outcome."""
        if outcome.status not in {"applied", "indeterminate"}:
            raise ValueError(f"Cannot complete upload journal record with status '{outcome.status}'.")
        if outcome.change_id not in self._records:
            raise ValueError("Cannot complete an upload journal record that was not prepared.")
        self._set_record(outcome, status=outcome.status)

    def discard(self, change_id: str) -> None:
        """Remove a pending record after Canvas definitively rejected a request."""
        if change_id not in self._records:
            return
        updated_records = dict(self._records)
        del updated_records[change_id]
        self._write_records(updated_records)
        self._records = updated_records

    def _set_record(self, outcome: GradebookUploadOutcome, *, status: str) -> None:
        updated_records = dict(self._records)
        updated_records[outcome.change_id] = {
            "status": status,
            "user_id": outcome.user_id,
            "has_submission_comment": outcome.has_submission_comment,
            "recorded_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        }
        self._write_records(updated_records)
        self._records = updated_records

    def _write_records(self, records: dict[str, dict[str, object]]) -> None:
        payload = {
            "schema_version": UPLOAD_JOURNAL_SCHEMA_VERSION,
            "records": records,
        }
        self.path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = self.path.with_name(f"{self.path.name}.tmp")
        temporary_path.write_text(
            json.dumps(payload, ensure_ascii=True, indent=2, sort_keys=True) + "\n",
            encoding="utf-8",
        )
        temporary_path.replace(self.path)


@dataclass(frozen=True)
class _MappedColumn:
    sheet_name: str
    column_name: str


@dataclass(frozen=True)
class _RatingChoice:
    criterion_id: str
    criterion_description: str
    rating_id: str
    selection_label: str
    descriptor: str
    points: object


@dataclass(frozen=True)
class _GradeInput:
    user_id: int
    grade: str | int | float | None
    submission_comment: str | None
    source_row: int


@dataclass(frozen=True)
class _RubricInput:
    user_id: int
    criterion_id: str
    assessment: dict[str, Any] | None
    source_row: int


@dataclass(frozen=True)
class _GradebookSnapshot:
    grades: dict[int, object]
    submissions: dict[int, _SubmissionSnapshot]
    rubric_assessments: dict[tuple[int, str], _RubricSnapshot]


def create_gradebook_template(
    client: CanvasClient,
    course_id: int,
    assignment_id: int,
    output: str | Path = ".",
) -> Path:
    """Create a controlled XLSX template for one Canvas assignment."""
    assignment = client.get_assignment(
        course_id=course_id,
        assignment_id=assignment_id,
        include_rubric=True,
    )
    _ensure_supported_assignment(assignment)
    rubric_criteria = _validated_rubric_criteria(assignment)
    submissions = client.list_submissions(
        course_id=course_id,
        assignment_id=assignment_id,
        include_user=True,
        include_rubric_assessment=True,
    )
    submissions_by_user = _submissions_by_user(submissions)
    rating_catalog = _build_rating_catalog(rubric_criteria)

    output_path = _resolve_template_output_path(
        output_path=Path(output),
        assignment_id=assignment_id,
        assignment_name=_assignment_name(assignment),
    )
    workbook = Workbook()
    instructions_sheet = workbook.active
    instructions_sheet.title = INSTRUCTIONS_SHEET
    _write_instructions_sheet(instructions_sheet, assignment)
    _write_grades_sheet(workbook, submissions_by_user)
    rubric_rows = _write_rubric_sheet(
        workbook=workbook,
        submissions_by_user=submissions_by_user,
        rubric_criteria=rubric_criteria,
        rating_catalog=rating_catalog,
    )
    _write_rubric_catalog_sheet(workbook, rating_catalog)
    _write_mapping_sheet(workbook)
    _write_metadata_sheet(
        workbook=workbook,
        course_id=course_id,
        assignment_id=assignment_id,
        assignment=assignment,
        rubric_criteria=rubric_criteria,
    )
    _write_snapshot_sheet(
        workbook=workbook,
        submissions_by_user=submissions_by_user,
        rubric_criteria=rubric_criteria,
    )
    choice_ranges = _write_rubric_lists_sheet(workbook, rating_catalog)
    _add_descriptor_validations(
        worksheet=workbook[RUBRIC_SHEET],
        rubric_rows=rubric_rows,
        choice_ranges=choice_ranges,
    )

    workbook[METADATA_SHEET].sheet_state = "hidden"
    workbook[SNAPSHOT_SHEET].sheet_state = "hidden"
    workbook[RUBRIC_LISTS_SHEET].sheet_state = "hidden"
    workbook.save(output_path)
    return output_path


def build_gradebook_plan(
    client: CanvasClient,
    course_id: int,
    assignment_id: int,
    input_path: str | Path,
) -> GradebookPlan:
    """Read and preflight a gradebook without making Canvas mutations."""
    workbook = _load_gradebook_workbook(Path(input_path))
    try:
        metadata = _read_metadata(workbook)
        _validate_metadata(metadata, course_id=course_id, assignment_id=assignment_id)
        mappings = _read_mapping(workbook)
        grade_sheet, grade_columns = _resolve_mapped_table(
            workbook,
            mappings,
            prefix="grades.",
            required_fields=GRADE_MAPPING_FIELDS,
        )
        rubric_sheet, rubric_columns = _resolve_mapped_table(
            workbook,
            mappings,
            prefix="rubric.",
            required_fields=RUBRIC_MAPPING_FIELDS,
        )
        snapshot = _read_snapshot(workbook)

        assignment = client.get_assignment(
            course_id=course_id,
            assignment_id=assignment_id,
            include_rubric=True,
        )
        _ensure_supported_assignment(assignment)
        rubric_criteria = _validated_rubric_criteria(assignment)
        expected_fingerprint = metadata["Rubric_Fingerprint"]
        actual_fingerprint = rubric_fingerprint(rubric_criteria)
        if expected_fingerprint != actual_fingerprint:
            raise GradebookValidationError(
                [
                    "The assignment rubric has changed since this template was created. "
                    "Generate a new template before uploading."
                ]
            )

        submissions = client.list_submissions(
            course_id=course_id,
            assignment_id=assignment_id,
            include_user=True,
            include_rubric_assessment=True,
        )
        submissions_by_user = _submissions_by_user(submissions)
        rating_catalog = _build_rating_catalog(rubric_criteria)
        rubric_criteria_by_alias = index_rubric_criteria(rubric_criteria)

        grade_inputs, grade_errors = _read_grade_inputs(grade_sheet, grade_columns)
        rubric_inputs, rubric_errors = _read_rubric_inputs(
            worksheet=rubric_sheet,
            columns=rubric_columns,
            rating_catalog=rating_catalog,
            criteria_by_alias=rubric_criteria_by_alias,
        )
        errors = [*grade_errors, *rubric_errors]
        _validate_assignment_grade_values(
            assignment=assignment,
            grade_inputs=grade_inputs,
            errors=errors,
        )
        _validate_input_students(
            grade_inputs=grade_inputs,
            rubric_inputs=rubric_inputs,
            submissions_by_user=submissions_by_user,
            errors=errors,
        )
        _validate_snapshot_conflicts(
            grade_inputs=grade_inputs,
            rubric_inputs=rubric_inputs,
            snapshot=snapshot,
            submissions_by_user=submissions_by_user,
            errors=errors,
        )
        if errors:
            raise GradebookValidationError(errors)

        rubric_updates_by_user: dict[int, dict[str, dict[str, Any]]] = {}
        for rubric_input in rubric_inputs:
            if rubric_input.assessment is None:
                continue
            rubric_updates_by_user.setdefault(rubric_input.user_id, {})[rubric_input.criterion_id] = (
                rubric_input.assessment
            )

        changes: list[GradebookChange] = []
        for grade_input in grade_inputs.values():
            rubric_assessment = rubric_updates_by_user.get(grade_input.user_id, {})
            if (
                grade_input.grade is None
                and grade_input.submission_comment is None
                and not rubric_assessment
            ):
                continue
            submission = submissions_by_user[grade_input.user_id]
            changes.append(
                GradebookChange(
                    user_id=grade_input.user_id,
                    student_name=_submission_student_name(submission),
                    posted_grade=grade_input.grade,
                    submission_comment=grade_input.submission_comment,
                    submission_comment_attempt=(
                        _parse_positive_id(submission.get("attempt"))
                        if grade_input.submission_comment is not None
                        else None
                    ),
                    rubric_assessment=rubric_assessment,
                    baseline=_build_change_baseline(
                        submission=submission,
                        posted_grade=grade_input.grade,
                        rubric_assessment=rubric_assessment,
                    ),
                )
            )

        return GradebookPlan(
            course_id=course_id,
            assignment_id=assignment_id,
            template_id=metadata["Template_ID"],
            total_students=len(grade_inputs),
            unchanged_students=len(grade_inputs) - len(changes),
            changes=tuple(changes),
        )
    finally:
        workbook.close()


def apply_gradebook_plan(
    client: CanvasClient,
    plan: GradebookPlan,
    *,
    execute: bool = False,
    progress_callback: Callable[[str], None] | None = None,
    upload_journal: GradebookUploadJournal | None = None,
    retry_indeterminate: bool = False,
) -> GradebookUploadResult:
    """Run a validated plan, or report the planned requests in dry-run mode."""
    if execute and any(change.baseline is None for change in plan.changes):
        raise ValueError("Gradebook plans must be created by build_gradebook_plan before execution.")
    if execute and plan.changes and upload_journal is None:
        raise ValueError("Executing a gradebook plan requires an upload journal.")

    applied = 0
    failed = 0
    conflicted = 0
    indeterminate = 0
    journal_skipped = 0
    journal_errors = 0
    failures: list[str] = []
    outcomes: list[GradebookUploadOutcome] = []
    excluded = (
        upload_journal.excluded_change_ids(retry_indeterminate=retry_indeterminate)
        if upload_journal is not None
        else set()
    )

    for change in plan.changes:
        description = describe_gradebook_change(change)
        student_label = _student_label(change.user_id, change.student_name)
        change_id = _gradebook_change_id(plan, change)
        if change_id in excluded:
            journal_skipped += 1
            _report(
                progress_callback,
                f"Skipped {student_label}: this exact change is already recorded in the upload journal.",
            )
            continue
        if not execute:
            _report(progress_callback, f"Would update {student_label}: {description}.")
            continue

        if _requires_prewrite_recheck(change):
            try:
                current_submission = client.get_submission(
                    course_id=plan.course_id,
                    assignment_id=plan.assignment_id,
                    user_id=change.user_id,
                    include_rubric_assessment=bool(change.rubric_assessment),
                )
            except CanvasAPIError as exc:
                failed += 1
                failure = f"{student_label}: could not recheck the current Canvas submission: {exc}"
                failures.append(failure)
                outcomes.append(
                    GradebookUploadOutcome(
                        change_id=change_id,
                        user_id=change.user_id,
                        status="failed",
                        has_submission_comment=change.submission_comment is not None,
                    )
                )
                _report(progress_callback, f"Failed to update {failure}")
                continue
            conflict = _prewrite_conflict(change, current_submission)
            if conflict is not None:
                conflicted += 1
                failure = f"{student_label}: skipped because {conflict}"
                failures.append(failure)
                outcomes.append(
                    GradebookUploadOutcome(
                        change_id=change_id,
                        user_id=change.user_id,
                        status="conflicted",
                        has_submission_comment=change.submission_comment is not None,
                    )
                )
                _report(progress_callback, failure)
                continue

        pending_outcome = GradebookUploadOutcome(
            change_id=change_id,
            user_id=change.user_id,
            status="pending",
            has_submission_comment=change.submission_comment is not None,
        )
        try:
            upload_journal.prepare(pending_outcome)
        except OSError as exc:
            journal_errors += 1
            failure = (
                f"{student_label}: could not record the pending upload request; "
                f"Canvas was not contacted: {exc}"
            )
            failures.append(failure)
            _report(progress_callback, f"Stopped upload because {failure}")
            break

        try:
            client.grade_submission(
                course_id=plan.course_id,
                assignment_id=plan.assignment_id,
                user_id=change.user_id,
                posted_grade=change.posted_grade,
                submission_comment=change.submission_comment,
                submission_comment_attempt=change.submission_comment_attempt,
                rubric_assessment=change.rubric_assessment or None,
            )
        except CanvasTransportError as exc:
            status = "indeterminate" if exc.request_may_have_succeeded else "failed"
            outcome = GradebookUploadOutcome(
                change_id=change_id,
                user_id=change.user_id,
                status=status,
                has_submission_comment=change.submission_comment is not None,
            )
            outcomes.append(outcome)
            if status == "indeterminate":
                indeterminate += 1
                failure = (
                    f"{student_label}: upload outcome is indeterminate because Canvas may have "
                    f"applied the request before the transport failure: {exc}"
                )
                try:
                    upload_journal.complete(outcome)
                except OSError as journal_exc:
                    journal_errors += 1
                    failures.append(
                        f"{student_label}: Canvas outcome is indeterminate and the upload journal "
                        f"could not be saved: {journal_exc}"
                    )
                    _report(
                        progress_callback,
                        "Stopped upload because an indeterminate outcome could not be recorded.",
                    )
                    break
                _report(progress_callback, f"Could not confirm update for {failure}")
            else:
                failed += 1
                failure = f"{student_label}: {exc}"
                try:
                    upload_journal.discard(change_id)
                except OSError as journal_exc:
                    journal_errors += 1
                    failures.append(
                        f"{student_label}: Canvas rejected the request and the pending upload "
                        f"journal record could not be removed: {journal_exc}"
                    )
                    _report(
                        progress_callback,
                        "Stopped upload because a rejected request could not be cleared from the journal.",
                    )
                    break
                _report(progress_callback, f"Failed to update {failure}")
            failures.append(failure)
        except CanvasAPIError as exc:
            if _canvas_error_is_indeterminate(exc):
                indeterminate += 1
                failure = (
                    f"{student_label}: upload outcome is indeterminate because Canvas returned "
                    f"a server error after receiving the request: {exc}"
                )
                outcome = GradebookUploadOutcome(
                    change_id=change_id,
                    user_id=change.user_id,
                    status="indeterminate",
                    has_submission_comment=change.submission_comment is not None,
                )
                outcomes.append(outcome)
                failures.append(failure)
                try:
                    upload_journal.complete(outcome)
                except OSError as journal_exc:
                    journal_errors += 1
                    failures.append(
                        f"{student_label}: Canvas outcome is indeterminate and the upload journal "
                        f"could not be saved: {journal_exc}"
                    )
                    _report(
                        progress_callback,
                        "Stopped upload because an indeterminate outcome could not be recorded.",
                    )
                    break
                _report(progress_callback, f"Could not confirm update for {failure}")
            else:
                failed += 1
                failure = f"{student_label}: {exc}"
                failures.append(failure)
                outcomes.append(
                    GradebookUploadOutcome(
                        change_id=change_id,
                        user_id=change.user_id,
                        status="failed",
                        has_submission_comment=change.submission_comment is not None,
                    )
                )
                try:
                    upload_journal.discard(change_id)
                except OSError as journal_exc:
                    journal_errors += 1
                    failures.append(
                        f"{student_label}: Canvas rejected the request and the pending upload "
                        f"journal record could not be removed: {journal_exc}"
                    )
                    _report(
                        progress_callback,
                        "Stopped upload because a rejected request could not be cleared from the journal.",
                    )
                    break
                _report(progress_callback, f"Failed to update {failure}")
        else:
            applied += 1
            outcome = GradebookUploadOutcome(
                change_id=change_id,
                user_id=change.user_id,
                status="applied",
                has_submission_comment=change.submission_comment is not None,
            )
            outcomes.append(outcome)
            try:
                upload_journal.complete(outcome)
            except OSError as journal_exc:
                journal_errors += 1
                failures.append(
                    f"{student_label}: Canvas updated the submission, but the upload journal "
                    f"could not be saved: {journal_exc}"
                )
                _report(
                    progress_callback,
                    "Stopped upload because a completed update could not be recorded in the journal.",
                )
                break
            _report(progress_callback, f"Updated {student_label}: {description}.")

    return GradebookUploadResult(
        planned=len(plan.changes),
        applied=applied,
        failed=failed,
        conflicted=conflicted,
        indeterminate=indeterminate,
        journal_skipped=journal_skipped,
        journal_errors=journal_errors,
        skipped=plan.unchanged_students + journal_skipped,
        failures=tuple(failures),
        outcomes=tuple(outcomes),
    )


def describe_gradebook_change(change: GradebookChange) -> str:
    """Return a privacy-preserving summary of one planned update."""
    parts: list[str] = []
    if change.posted_grade is not None:
        parts.append("grade")
    if change.submission_comment is not None:
        parts.append("new submission comment")
    if change.rubric_assessment:
        count = len(change.rubric_assessment)
        noun = "criterion" if count == 1 else "criteria"
        parts.append(f"{count} rubric {noun}")
    return ", ".join(parts)


def default_gradebook_upload_journal_path(input_path: str | Path) -> Path:
    """Return the sidecar journal path used to prevent accidental replays."""
    workbook_path = Path(input_path)
    return workbook_path.with_name(f"{workbook_path.name}.upload-journal.json")


def default_gradebook_upload_lock_path(journal_path: str | Path) -> Path:
    """Return the advisory lock path paired with an upload journal."""
    path = Path(journal_path)
    return path.with_name(f"{path.name}.lock")


def _build_change_baseline(
    *,
    submission: dict[str, Any],
    posted_grade: str | int | float | None,
    rubric_assessment: dict[str, dict[str, Any]],
) -> _ChangeBaseline:
    rubric_baselines = {
        criterion_id: _rubric_snapshot_from_selection(
            _submission_rubric_selection(submission, criterion_id)
        )
        for criterion_id in rubric_assessment
    }
    return _ChangeBaseline(
        grade=_submission_grade(submission) if posted_grade is not None else None,
        rubric_assessments=rubric_baselines,
        submission=_submission_snapshot_from_submission(submission),
    )


def _requires_prewrite_recheck(change: GradebookChange) -> bool:
    return (
        change.posted_grade is not None
        or change.submission_comment is not None
        or bool(change.rubric_assessment)
    )


def _prewrite_conflict(
    change: GradebookChange,
    current_submission: dict[str, Any],
) -> str | None:
    baseline = change.baseline
    if baseline is None:
        return "the plan was not created by gradebook validation"
    if change.posted_grade is not None and not _same_canvas_value(
        baseline.grade,
        _submission_grade(current_submission),
    ):
        return "the grade changed in Canvas after preflight"
    if (
        (change.submission_comment is not None or bool(change.rubric_assessment))
        and not _same_submission_state(
            baseline.submission,
            current_submission,
        )
    ):
        return "the submission changed in Canvas after preflight"
    for criterion_id in change.rubric_assessment:
        expected = baseline.rubric_assessments.get(criterion_id)
        if expected is None:
            return f"the baseline is missing rubric criterion '{criterion_id}'"
        current = _rubric_snapshot_from_selection(
            _submission_rubric_selection(current_submission, criterion_id)
        )
        if not (
            _same_canvas_value(expected.rating_id, current.rating_id)
            and _same_canvas_value(expected.points, current.points)
            and _same_canvas_value(expected.comment, current.comment)
        ):
            return f"rubric criterion '{criterion_id}' changed in Canvas after preflight"
    return None


def _rubric_snapshot_from_selection(selection: dict[str, Any]) -> _RubricSnapshot:
    return _RubricSnapshot(
        rating_id=selection.get("rating_id"),
        points=selection.get("points"),
        comment=selection.get("comments"),
    )


def _submission_snapshot_from_submission(submission: dict[str, Any]) -> _SubmissionSnapshot:
    return _SubmissionSnapshot(
        attempt=submission.get("attempt"),
        submitted_at=submission.get("submitted_at"),
    )


def _same_submission_state(
    expected: _SubmissionSnapshot,
    current_submission: dict[str, Any],
) -> bool:
    current = _submission_snapshot_from_submission(current_submission)
    return (
        _same_canvas_value(expected.attempt, current.attempt)
        and _same_canvas_value(expected.submitted_at, current.submitted_at)
    )


def _gradebook_change_id(plan: GradebookPlan, change: GradebookChange) -> str:
    payload = {
        "assignment_id": plan.assignment_id,
        "course_id": plan.course_id,
        "posted_grade": change.posted_grade,
        "rubric_assessment": {
            criterion_id: {
                field: assessment.get(field)
                for field in ("rating_id", "points", "comments")
                if assessment.get(field) is not None
            }
            for criterion_id, assessment in sorted(change.rubric_assessment.items())
        },
        "submission_comment": change.submission_comment,
        "submission_comment_attempt": change.submission_comment_attempt,
        "template_id": plan.template_id,
        "user_id": change.user_id,
    }
    encoded = json.dumps(
        payload,
        ensure_ascii=True,
        separators=(",", ":"),
        sort_keys=True,
        allow_nan=False,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _canvas_error_is_indeterminate(error: CanvasAPIError) -> bool:
    return error.status_code is not None and error.status_code >= 500


def rubric_fingerprint(rubric_criteria: list[dict[str, Any]]) -> str:
    """Return a stable fingerprint for the relevant rubric definition fields."""
    normalized_criteria: list[dict[str, object]] = []
    for criterion in rubric_criteria:
        criterion_id = identifier_or_none(criterion.get("id"))
        if criterion_id is None:
            continue
        normalized_ratings: list[dict[str, object]] = []
        for rating in rubric_ratings(criterion):
            rating_id = identifier_or_none(rating.get("id"))
            if rating_id is None:
                continue
            normalized_ratings.append(
                {
                    "id": rating_id,
                    "description": rubric_rating_description(rating) or "",
                    "points": _normalise_fingerprint_value(rating.get("points")),
                }
            )
        normalized_criteria.append(
            {
                "id": criterion_id,
                "description": rubric_criterion_description(criterion) or "",
                "ratings": sorted(normalized_ratings, key=lambda rating: str(rating["id"])),
            }
        )
    encoded = json.dumps(
        sorted(normalized_criteria, key=lambda criterion: str(criterion["id"])),
        ensure_ascii=True,
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _ensure_supported_assignment(assignment: dict[str, Any]) -> None:
    if bool(assignment.get("anonymous_grading")):
        raise ValueError("Gradebook upload does not support anonymously graded assignments.")
    if bool(assignment.get("moderated_grading")):
        raise ValueError("Gradebook upload does not support moderated assignments.")
    if assignment.get("group_category_id") not in (None, "", 0):
        raise ValueError("Gradebook upload does not support group assignments.")


def _validate_assignment_grade_values(
    *,
    assignment: dict[str, Any],
    grade_inputs: dict[int, _GradeInput],
    errors: list[str],
) -> None:
    grading_type = (_text_or_none(assignment.get("grading_type")) or "").casefold()
    changed_grades = [
        grade_input for grade_input in grade_inputs.values() if grade_input.grade is not None
    ]
    if not changed_grades:
        return
    if grading_type == "not_graded":
        for grade_input in changed_grades:
            errors.append(
                f"Grades row {grade_input.source_row} cannot set a grade because this "
                "Canvas assignment is not graded."
            )
        return
    if grading_type != "pass_fail":
        return

    points_possible = _parse_optional_points(
        assignment.get("points_possible"),
        location="Canvas assignment points_possible",
        errors=errors,
    )
    for grade_input in changed_grades:
        if not _is_valid_pass_fail_grade(grade_input.grade, points_possible):
            errors.append(
                f"Grades row {grade_input.source_row} must use pass, complete, fail, "
                "incomplete, 0%, 100%, 0, or the assignment's points possible for a "
                "pass/fail assignment."
            )


def _is_valid_pass_fail_grade(
    value: str | int | float,
    points_possible: int | float | None,
) -> bool:
    if isinstance(value, str):
        normalized = value.strip().casefold()
        if normalized in {"pass", "complete", "fail", "incomplete"}:
            return True
        if normalized.endswith("%"):
            try:
                percentage = Decimal(normalized.removesuffix("%").strip())
            except InvalidOperation:
                return False
            return percentage in {Decimal("0"), Decimal("100")}
    return _same_canvas_value(value, 0) or (
        points_possible is not None and _same_canvas_value(value, points_possible)
    )


def _validated_rubric_criteria(assignment: dict[str, Any]) -> list[dict[str, Any]]:
    rubric_criteria = extract_rubric_criteria(assignment)
    seen_ids: set[str] = set()
    errors: list[str] = []
    for index, criterion in enumerate(rubric_criteria, start=1):
        criterion_id = identifier_or_none(criterion.get("id"))
        if criterion_id is None:
            errors.append(f"Canvas rubric criterion {index} does not have an identifier.")
            continue
        if criterion_id in seen_ids:
            errors.append(f"Canvas rubric contains duplicate criterion ID '{criterion_id}'.")
        seen_ids.add(criterion_id)
    if errors:
        raise ValueError("\n".join(errors))
    return rubric_criteria


def _submissions_by_user(submissions: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    records: dict[int, dict[str, Any]] = {}
    errors: list[str] = []
    for index, submission in enumerate(submissions, start=1):
        user_id = _parse_positive_id(submission.get("user_id"))
        if user_id is None:
            errors.append(f"Canvas submission {index} does not contain a positive user ID.")
            continue
        if user_id in records:
            errors.append(f"Canvas returned multiple submission records for user ID {user_id}.")
            continue
        records[user_id] = submission
    if errors:
        raise ValueError("\n".join(errors))
    return records


def _build_rating_catalog(
    rubric_criteria: list[dict[str, Any]],
) -> dict[str, list[_RatingChoice]]:
    catalog: dict[str, list[_RatingChoice]] = {}
    for criterion in rubric_criteria:
        criterion_id = identifier_or_none(criterion.get("id"))
        if criterion_id is None:
            continue
        criterion_description = rubric_criterion_description(criterion) or "(No criterion description)"
        raw_choices: list[tuple[str, str, object, str]] = []
        for rating in rubric_ratings(criterion):
            rating_id = identifier_or_none(rating.get("id"))
            if rating_id is None:
                continue
            descriptor = rubric_rating_description(rating) or "(No descriptor)"
            base_label = f"{_format_points(rating.get('points'))} - {descriptor}"
            raw_choices.append((rating_id, descriptor, rating.get("points"), base_label))

        labels = [choice[3] for choice in raw_choices]
        choices: list[_RatingChoice] = []
        for rating_id, descriptor, points, base_label in raw_choices:
            selection_label = (
                f"{base_label} [{rating_id}]" if labels.count(base_label) > 1 else base_label
            )
            choices.append(
                _RatingChoice(
                    criterion_id=criterion_id,
                    criterion_description=criterion_description,
                    rating_id=rating_id,
                    selection_label=selection_label,
                    descriptor=descriptor,
                    points=points,
                )
            )
        catalog[criterion_id] = choices
    return catalog


def _resolve_template_output_path(
    *,
    output_path: Path,
    assignment_id: int,
    assignment_name: str,
) -> Path:
    file_name = (
        f"canvas_gradebook_{assignment_id}_{_slugify_filename_part(assignment_name)}_"
        f"{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    )
    if output_path.suffix.lower() == ".xlsx":
        output_path.parent.mkdir(parents=True, exist_ok=True)
        return output_path

    output_path.mkdir(parents=True, exist_ok=True)
    return output_path / file_name


def _write_instructions_sheet(worksheet: Any, assignment: dict[str, Any]) -> None:
    grading_type = _text_or_none(assignment.get("grading_type")) or "unknown"
    points_possible = _text_or_none(assignment.get("points_possible")) or "unknown"
    instructions = [
        ("Canvas gradebook template",),
        (f"Assignment grading type: {grading_type}; points possible: {points_possible}.",),
        (
            "Use this workbook only for the course and assignment recorded in Canvas Metadata.",
        ),
        (
            "Grades: set Grade_Action to SET and enter New_Grade. KEEP leaves the current grade unchanged.",
        ),
        (
            "Comments: set Comment_Action to ADD and enter New_Submission_Comment. "
            "Canvas adds a new comment; it does not replace existing feedback.",
        ),
        (
            "Rubric: set Action to SET. Choose a descriptor from the dropdown or enter "
            "Selected_Rating_ID, then optionally override Points and add a criterion comment.",
        ),
        (
            "Blank editable values are only valid with KEEP. The uploader rejects ambiguous rows "
            "instead of silently changing or clearing marks.",
        ),
        (
            "If you rename data columns, update Canvas Mapping with the new sheet and column names.",
        ),
        (
            "Run gradebook validate before gradebook upload. Upload is a dry run until --execute is supplied.",
        ),
    ]
    for row in instructions:
        worksheet.append(row)
    worksheet["A1"].font = Font(bold=True, size=14)
    worksheet.column_dimensions["A"].width = 120
    worksheet.sheet_view.showGridLines = False


def _sanitize_cell_value(val: Any) -> Any:
    if isinstance(val, str) and val.startswith("="):
        return f"'{val}"
    return val


def _append_sheet_row(worksheet: Any, row: Iterable[Any]) -> None:
    worksheet.append([_sanitize_cell_value(cell) for cell in row])


def _write_grades_sheet(
    workbook: Workbook,
    submissions_by_user: dict[int, dict[str, Any]],
) -> None:
    worksheet = workbook.create_sheet(title=GRADES_SHEET)
    worksheet.append(GRADE_HEADERS)
    for user_id, submission in submissions_by_user.items():
        user = _submission_user(submission)
        _append_sheet_row(
            worksheet,
            [
                user_id,
                user.get("name"),
                user.get("login_id"),
                _submission_grade(submission),
                None,
                "KEEP",
                None,
                "KEEP",
            ],
        )
    _format_table(worksheet)
    _shade_read_only_columns(worksheet, ("A", "B", "C", "D"))


def _write_rubric_sheet(
    *,
    workbook: Workbook,
    submissions_by_user: dict[int, dict[str, Any]],
    rubric_criteria: list[dict[str, Any]],
    rating_catalog: dict[str, list[_RatingChoice]],
) -> dict[str, list[int]]:
    worksheet = workbook.create_sheet(title=RUBRIC_SHEET)
    worksheet.append(RUBRIC_HEADERS)
    rows_by_criterion: dict[str, list[int]] = {}
    for user_id, submission in submissions_by_user.items():
        user = _submission_user(submission)
        for criterion in rubric_criteria:
            criterion_id = identifier_or_none(criterion.get("id"))
            if criterion_id is None:
                continue
            selection = _submission_rubric_selection(submission, criterion_id)
            current_rating_id = identifier_or_none(selection.get("rating_id"))
            current_rating = find_rubric_rating(criterion, current_rating_id)
            choice = _rating_choice_by_id(rating_catalog.get(criterion_id, []), current_rating_id)
            current_descriptor = (
                choice.selection_label
                if choice is not None
                else rubric_rating_description(current_rating)
                or _text_or_none(selection.get("description"))
            )
            _append_sheet_row(
                worksheet,
                [
                    user_id,
                    user.get("name"),
                    criterion_id,
                    rubric_criterion_description(criterion),
                    current_descriptor,
                    selection.get("points"),
                    selection.get("comments"),
                    None,
                    None,
                    None,
                    None,
                    "KEEP",
                ],
            )
            rows_by_criterion.setdefault(criterion_id, []).append(worksheet.max_row)
    _format_table(worksheet)
    _shade_read_only_columns(worksheet, ("A", "B", "C", "D", "E", "F", "G"))
    worksheet.column_dimensions["I"].hidden = True
    return rows_by_criterion


def _write_rubric_catalog_sheet(
    workbook: Workbook,
    rating_catalog: dict[str, list[_RatingChoice]],
) -> None:
    worksheet = workbook.create_sheet(title=RUBRIC_CATALOG_SHEET)
    worksheet.append(RUBRIC_CATALOG_HEADERS)
    for choices in rating_catalog.values():
        for choice in choices:
            _append_sheet_row(
                worksheet,
                [
                    choice.criterion_id,
                    choice.criterion_description,
                    choice.rating_id,
                    choice.selection_label,
                    choice.descriptor,
                    choice.points,
                ],
            )
    _format_table(worksheet)
    _shade_read_only_columns(worksheet, ("A", "B", "C", "D", "E", "F"))


def _write_mapping_sheet(workbook: Workbook) -> None:
    worksheet = workbook.create_sheet(title=MAPPING_SHEET)
    worksheet.append(MAPPING_HEADERS)
    rows = [
        ("grades.student_id", GRADES_SHEET, "Canvas_User_ID"),
        ("grades.new_grade", GRADES_SHEET, "New_Grade"),
        ("grades.grade_action", GRADES_SHEET, "Grade_Action"),
        ("grades.submission_comment", GRADES_SHEET, "New_Submission_Comment"),
        ("grades.comment_action", GRADES_SHEET, "Comment_Action"),
        ("rubric.student_id", RUBRIC_SHEET, "Canvas_User_ID"),
        ("rubric.criterion_id", RUBRIC_SHEET, "Criterion_ID"),
        ("rubric.selected_descriptor", RUBRIC_SHEET, "Selected_Descriptor"),
        ("rubric.selected_rating_id", RUBRIC_SHEET, "Selected_Rating_ID"),
        ("rubric.points", RUBRIC_SHEET, "Points"),
        ("rubric.criterion_comment", RUBRIC_SHEET, "Criterion_Comment"),
        ("rubric.action", RUBRIC_SHEET, "Action"),
    ]
    for row in rows:
        _append_sheet_row(worksheet, row)
    _format_table(worksheet)


def _write_metadata_sheet(
    *,
    workbook: Workbook,
    course_id: int,
    assignment_id: int,
    assignment: dict[str, Any],
    rubric_criteria: list[dict[str, Any]],
) -> None:
    worksheet = workbook.create_sheet(title=METADATA_SHEET)
    worksheet.append(METADATA_HEADERS)
    for row in (
        ("Schema_Version", GRADEBOOK_SCHEMA_VERSION),
        ("Template_ID", str(uuid.uuid4())),
        ("Course_ID", course_id),
        ("Assignment_ID", assignment_id),
        ("Assignment_Name", _assignment_name(assignment)),
        ("Grading_Type", _text_or_none(assignment.get("grading_type")) or "unknown"),
        ("Points_Possible", _text_or_none(assignment.get("points_possible")) or "unknown"),
        ("Rubric_Fingerprint", rubric_fingerprint(rubric_criteria)),
        ("Template_Created_At", datetime.now().astimezone().isoformat(timespec="seconds")),
    ):
        _append_sheet_row(worksheet, row)
    _format_table(worksheet)


def _write_snapshot_sheet(
    *,
    workbook: Workbook,
    submissions_by_user: dict[int, dict[str, Any]],
    rubric_criteria: list[dict[str, Any]],
) -> None:
    worksheet = workbook.create_sheet(title=SNAPSHOT_SHEET)
    worksheet.append(SNAPSHOT_HEADERS)
    for user_id, submission in submissions_by_user.items():
        _append_sheet_row(
            worksheet,
            (
                "grade",
                user_id,
                None,
                submission.get("attempt"),
                submission.get("submitted_at"),
                _submission_grade(submission),
                None,
                None,
                None,
            ),
        )
        for criterion in rubric_criteria:
            criterion_id = identifier_or_none(criterion.get("id"))
            if criterion_id is None:
                continue
            selection = _submission_rubric_selection(submission, criterion_id)
            _append_sheet_row(
                worksheet,
                (
                    "rubric",
                    user_id,
                    criterion_id,
                    submission.get("attempt"),
                    submission.get("submitted_at"),
                    None,
                    selection.get("rating_id"),
                    selection.get("points"),
                    selection.get("comments"),
                ),
            )
    _format_table(worksheet)


def _write_rubric_lists_sheet(
    workbook: Workbook,
    rating_catalog: dict[str, list[_RatingChoice]],
) -> dict[str, str]:
    worksheet = workbook.create_sheet(title=RUBRIC_LISTS_SHEET)
    choice_ranges: dict[str, str] = {}
    for column_number, (criterion_id, choices) in enumerate(rating_catalog.items(), start=1):
        if not choices:
            continue
        worksheet.cell(row=1, column=column_number, value=criterion_id)
        for row_number, choice in enumerate(choices, start=2):
            worksheet.cell(
                row=row_number,
                column=column_number,
                value=_sanitize_cell_value(choice.selection_label),
            )
        range_name = f"RubricChoices{column_number}"
        column_letter = get_column_letter(column_number)
        reference = f"'{RUBRIC_LISTS_SHEET}'!${column_letter}$2:${column_letter}${len(choices) + 1}"
        workbook.defined_names.add(DefinedName(range_name, attr_text=reference))
        choice_ranges[criterion_id] = range_name
    return choice_ranges


def _add_descriptor_validations(
    *,
    worksheet: Any,
    rubric_rows: dict[str, list[int]],
    choice_ranges: dict[str, str],
) -> None:
    descriptor_column = get_column_letter(RUBRIC_HEADERS.index("Selected_Descriptor") + 1)
    for criterion_id, row_numbers in rubric_rows.items():
        range_name = choice_ranges.get(criterion_id)
        if range_name is None:
            continue
        validation = DataValidation(type="list", formula1=f"={range_name}", allow_blank=True)
        validation.error = "Choose a descriptor from the list, or use the hidden Selected_Rating_ID column."
        validation.errorTitle = "Unknown rubric descriptor"
        worksheet.add_data_validation(validation)
        for row_number in row_numbers:
            validation.add(f"{descriptor_column}{row_number}")


def _format_table(worksheet: Any) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(column_cells[0].column)
        widest = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_letter].width = min(max(widest + 2, 12), 48)


def _shade_read_only_columns(worksheet: Any, columns: tuple[str, ...]) -> None:
    fill = PatternFill(fill_type="solid", fgColor="E7E6E6")
    for column_letter in columns:
        for cell in worksheet[column_letter]:
            if cell.row > 1:
                cell.fill = fill


def _load_gradebook_workbook(input_path: Path) -> Workbook:
    if not input_path.is_file():
        raise ValueError(f"Gradebook input file does not exist: {input_path}")
    try:
        return load_workbook(input_path, data_only=False)
    except (BadZipFile, InvalidFileException, KeyError, OSError, ParseError, ValueError) as exc:
        raise ValueError(f"Could not read gradebook workbook '{input_path}': {exc}") from exc


def _read_metadata(workbook: Workbook) -> dict[str, str]:
    worksheet = _required_sheet(workbook, METADATA_SHEET)
    headers = _header_index(worksheet)
    key_column = _required_header(headers, "Key", METADATA_SHEET)
    value_column = _required_header(headers, "Value", METADATA_SHEET)
    metadata: dict[str, str] = {}
    errors: list[str] = []
    for row_number in range(2, worksheet.max_row + 1):
        key_cell = worksheet.cell(row=row_number, column=key_column)
        value_cell = worksheet.cell(row=row_number, column=value_column)
        if _is_blank(key_cell.value) and _is_blank(value_cell.value):
            continue
        if key_cell.data_type == "f" or value_cell.data_type == "f":
            errors.append(f"{METADATA_SHEET}!{row_number} cannot contain formulas.")
            continue
        key = _text_or_none(key_cell.value)
        value = _text_or_none(value_cell.value)
        if key is None or value is None:
            errors.append(f"{METADATA_SHEET}!{row_number} must contain a key and value.")
            continue
        if key in metadata:
            errors.append(f"{METADATA_SHEET}!{row_number} duplicates metadata key '{key}'.")
            continue
        metadata[key] = value
    required_keys = (
        "Schema_Version",
        "Template_ID",
        "Course_ID",
        "Assignment_ID",
        "Rubric_Fingerprint",
    )
    for key in required_keys:
        if key not in metadata:
            errors.append(f"{METADATA_SHEET} is missing required metadata key '{key}'.")
    if errors:
        raise GradebookValidationError(errors)
    return metadata


def _validate_metadata(metadata: dict[str, str], *, course_id: int, assignment_id: int) -> None:
    errors: list[str] = []
    if metadata["Schema_Version"] != GRADEBOOK_SCHEMA_VERSION:
        errors.append(
            f"Unsupported gradebook schema version '{metadata['Schema_Version']}'. "
            f"Expected '{GRADEBOOK_SCHEMA_VERSION}'."
        )
    try:
        template_id = uuid.UUID(metadata["Template_ID"])
    except ValueError:
        errors.append(f"Workbook Template_ID '{metadata['Template_ID']}' is not a valid UUID.")
    else:
        if str(template_id) != metadata["Template_ID"].casefold():
            errors.append(f"Workbook Template_ID '{metadata['Template_ID']}' is not canonical.")
    if _parse_positive_id(metadata["Course_ID"]) != course_id:
        errors.append(
            f"Workbook Course_ID '{metadata['Course_ID']}' does not match --course {course_id}."
        )
    if _parse_positive_id(metadata["Assignment_ID"]) != assignment_id:
        errors.append(
            f"Workbook Assignment_ID '{metadata['Assignment_ID']}' does not match --assignment {assignment_id}."
        )
    if errors:
        raise GradebookValidationError(errors)


def _read_mapping(workbook: Workbook) -> dict[str, _MappedColumn]:
    worksheet = _required_sheet(workbook, MAPPING_SHEET)
    headers = _header_index(worksheet)
    field_column = _required_header(headers, "Field", MAPPING_SHEET)
    sheet_column = _required_header(headers, "Sheet", MAPPING_SHEET)
    column_column = _required_header(headers, "Column", MAPPING_SHEET)
    mappings: dict[str, _MappedColumn] = {}
    errors: list[str] = []
    for row_number in range(2, worksheet.max_row + 1):
        cells = (
            worksheet.cell(row=row_number, column=field_column),
            worksheet.cell(row=row_number, column=sheet_column),
            worksheet.cell(row=row_number, column=column_column),
        )
        if all(_is_blank(cell.value) for cell in cells):
            continue
        if any(cell.data_type == "f" for cell in cells):
            errors.append(f"{MAPPING_SHEET}!{row_number} cannot contain formulas.")
            continue
        field_name = _text_or_none(cells[0].value)
        sheet_name = _text_or_none(cells[1].value)
        column_name = _text_or_none(cells[2].value)
        if field_name is None or sheet_name is None or column_name is None:
            errors.append(f"{MAPPING_SHEET}!{row_number} must contain Field, Sheet, and Column.")
            continue
        normalized_field = field_name.casefold()
        if normalized_field not in ALL_MAPPING_FIELDS:
            errors.append(f"{MAPPING_SHEET}!{row_number} has unsupported field '{field_name}'.")
            continue
        if normalized_field in mappings:
            errors.append(f"{MAPPING_SHEET}!{row_number} duplicates field '{field_name}'.")
            continue
        mappings[normalized_field] = _MappedColumn(sheet_name=sheet_name, column_name=column_name)
    if errors:
        raise GradebookValidationError(errors)
    return mappings


def _resolve_mapped_table(
    workbook: Workbook,
    mappings: dict[str, _MappedColumn],
    *,
    prefix: str,
    required_fields: tuple[str, ...],
) -> tuple[Any, dict[str, int]]:
    errors: list[str] = []
    missing = [field for field in required_fields if field not in mappings]
    for field in missing:
        errors.append(f"{MAPPING_SHEET} is missing required field '{field}'.")
    if errors:
        raise GradebookValidationError(errors)

    table_mappings = {
        field: mapping for field, mapping in mappings.items() if field.startswith(prefix)
    }
    sheet_names = {mapping.sheet_name for mapping in table_mappings.values()}
    if len(sheet_names) != 1:
        raise GradebookValidationError(
            [f"All {prefix.removesuffix('.')} mapping fields must reference the same worksheet."]
        )
    sheet_name = next(iter(sheet_names))
    worksheet = _required_sheet(workbook, sheet_name)
    headers = _header_index(worksheet)
    columns: dict[str, int] = {}
    seen_columns: dict[int, str] = {}
    for field, mapping in table_mappings.items():
        column_number = headers.get(mapping.column_name.casefold())
        if column_number is None:
            errors.append(
                f"{MAPPING_SHEET} maps '{field}' to missing column "
                f"'{mapping.column_name}' on sheet '{sheet_name}'."
            )
            continue
        if column_number in seen_columns:
            errors.append(
                f"{MAPPING_SHEET} maps multiple fields ('{seen_columns[column_number]}' and '{field}') "
                f"to the same column '{mapping.column_name}' on sheet '{sheet_name}'."
            )
            continue
        seen_columns[column_number] = field
        columns[field] = column_number
    if errors:
        raise GradebookValidationError(errors)
    return worksheet, columns


def _read_snapshot(workbook: Workbook) -> _GradebookSnapshot:
    worksheet = _required_sheet(workbook, SNAPSHOT_SHEET)
    headers = _header_index(worksheet)
    header_columns = {
        header: _required_header(headers, header, SNAPSHOT_SHEET) for header in SNAPSHOT_HEADERS
    }
    grades: dict[int, object] = {}
    submissions: dict[int, _SubmissionSnapshot] = {}
    rubric_assessments: dict[tuple[int, str], _RubricSnapshot] = {}
    errors: list[str] = []
    for row_number in range(2, worksheet.max_row + 1):
        values = {
            header: worksheet.cell(row=row_number, column=column)
            for header, column in header_columns.items()
        }
        if all(_is_blank(cell.value) for cell in values.values()):
            continue
        if any(cell.data_type == "f" for cell in values.values()):
            errors.append(f"{SNAPSHOT_SHEET}!{row_number} cannot contain formulas.")
            continue
        record_type = (_text_or_none(values["Record_Type"].value) or "").casefold()
        user_id = _parse_positive_id(values["Canvas_User_ID"].value)
        if user_id is None:
            errors.append(f"{SNAPSHOT_SHEET}!{row_number} has an invalid Canvas_User_ID.")
            continue
        if record_type == "grade":
            if user_id in grades:
                errors.append(f"{SNAPSHOT_SHEET}!{row_number} duplicates the grade snapshot for {user_id}.")
                continue
            grades[user_id] = values["Grade"].value
            submissions[user_id] = _SubmissionSnapshot(
                attempt=values["Submission_Attempt"].value,
                submitted_at=values["Submitted_At"].value,
            )
            continue
        if record_type == "rubric":
            criterion_id = identifier_or_none(values["Criterion_ID"].value)
            if criterion_id is None:
                errors.append(f"{SNAPSHOT_SHEET}!{row_number} has an invalid Criterion_ID.")
                continue
            key = (user_id, criterion_id)
            if key in rubric_assessments:
                errors.append(
                    f"{SNAPSHOT_SHEET}!{row_number} duplicates the rubric snapshot for "
                    f"user {user_id}, criterion '{criterion_id}'."
                )
                continue
            rubric_assessments[key] = _RubricSnapshot(
                rating_id=values["Rating_ID"].value,
                points=values["Points"].value,
                comment=values["Criterion_Comment"].value,
            )
            continue
        errors.append(
            f"{SNAPSHOT_SHEET}!{row_number} has unsupported Record_Type "
            f"'{values['Record_Type'].value}'."
        )
    if errors:
        raise GradebookValidationError(errors)
    return _GradebookSnapshot(
        grades=grades,
        submissions=submissions,
        rubric_assessments=rubric_assessments,
    )


def _read_grade_inputs(
    worksheet: Any,
    columns: dict[str, int],
) -> tuple[dict[int, _GradeInput], list[str]]:
    inputs: dict[int, _GradeInput] = {}
    errors: list[str] = []
    fields = GRADE_MAPPING_FIELDS
    for row_number in range(2, worksheet.max_row + 1):
        values = _read_mapped_values(worksheet, row_number, columns, fields, errors)
        if all(_is_blank(value) for value in values.values()):
            continue
        user_id = _parse_positive_id(values["grades.student_id"])
        if user_id is None:
            errors.append(
                f"{worksheet.title}!{row_number} has an invalid Canvas user ID in "
                f"'{_mapped_header_name(worksheet, columns, 'grades.student_id')}'."
            )
            continue
        if user_id in inputs:
            errors.append(f"{worksheet.title}!{row_number} duplicates Canvas user ID {user_id}.")
            continue

        grade_action = _parse_action(
            values["grades.grade_action"],
            allowed=("KEEP", "SET"),
            location=f"{worksheet.title}!{row_number} Grade_Action",
            errors=errors,
        )
        comment_action = _parse_action(
            values["grades.comment_action"],
            allowed=("KEEP", "ADD"),
            location=f"{worksheet.title}!{row_number} Comment_Action",
            errors=errors,
        )
        grade_value: str | int | float | None = None
        if grade_action == "SET":
            grade_value = _parse_grade_value(
                values["grades.new_grade"],
                location=f"{worksheet.title}!{row_number} New_Grade",
                errors=errors,
            )
        elif not _is_blank(values["grades.new_grade"]):
            errors.append(
                f"{worksheet.title}!{row_number} has New_Grade but Grade_Action is KEEP."
            )

        comment_value = _text_or_none(values["grades.submission_comment"])
        if comment_action == "ADD":
            if comment_value is None:
                errors.append(
                    f"{worksheet.title}!{row_number} has Comment_Action ADD without a "
                    "New_Submission_Comment."
                )
        elif comment_value is not None:
            errors.append(
                f"{worksheet.title}!{row_number} has New_Submission_Comment but Comment_Action is KEEP."
            )

        inputs[user_id] = _GradeInput(
            user_id=user_id,
            grade=grade_value,
            submission_comment=comment_value if comment_action == "ADD" else None,
            source_row=row_number,
        )
    return inputs, errors


def _read_rubric_inputs(
    *,
    worksheet: Any,
    columns: dict[str, int],
    rating_catalog: dict[str, list[_RatingChoice]],
    criteria_by_alias: dict[str, dict[str, Any]],
) -> tuple[list[_RubricInput], list[str]]:
    inputs: list[_RubricInput] = []
    errors: list[str] = []
    fields = (*RUBRIC_MAPPING_FIELDS, *OPTIONAL_RUBRIC_MAPPING_FIELDS)
    seen_records: set[tuple[int, str]] = set()
    for row_number in range(2, worksheet.max_row + 1):
        values = _read_mapped_values(worksheet, row_number, columns, fields, errors)
        if all(_is_blank(value) for value in values.values()):
            continue
        user_id = _parse_positive_id(values["rubric.student_id"])
        if user_id is None:
            errors.append(f"{worksheet.title}!{row_number} has an invalid Canvas user ID.")
            continue
        submitted_criterion_id = identifier_or_none(values["rubric.criterion_id"])
        criterion = (
            criteria_by_alias.get(submitted_criterion_id)
            if submitted_criterion_id is not None
            else None
        )
        if criterion is None:
            errors.append(
                f"{worksheet.title}!{row_number} references unknown criterion "
                f"'{values['rubric.criterion_id']}'."
            )
            continue
        criterion_id = identifier_or_none(criterion.get("id"))
        if criterion_id is None:
            errors.append(f"{worksheet.title}!{row_number} could not resolve a Canvas criterion ID.")
            continue
        record_key = (user_id, criterion_id)
        if record_key in seen_records:
            errors.append(
                f"{worksheet.title}!{row_number} duplicates rubric criterion '{criterion_id}' "
                f"for Canvas user ID {user_id}."
            )
            continue
        seen_records.add(record_key)

        action = _parse_action(
            values["rubric.action"],
            allowed=("KEEP", "SET"),
            location=f"{worksheet.title}!{row_number} Action",
            errors=errors,
        )
        descriptor = _text_or_none(values["rubric.selected_descriptor"])
        selected_rating_id = identifier_or_none(values.get("rubric.selected_rating_id"))
        points = _parse_optional_points(
            values["rubric.points"],
            location=f"{worksheet.title}!{row_number} Points",
            errors=errors,
        )
        comment = _text_or_none(values["rubric.criterion_comment"])
        if action == "KEEP":
            if any(value is not None for value in (descriptor, selected_rating_id, points, comment)):
                errors.append(
                    f"{worksheet.title}!{row_number} has rubric values but Action is KEEP."
                )
            inputs.append(
                _RubricInput(
                    user_id=user_id,
                    criterion_id=criterion_id,
                    assessment=None,
                    source_row=row_number,
                )
            )
            continue

        choices = rating_catalog.get(criterion_id, [])
        choice_from_descriptor = _rating_choice_by_label(choices, descriptor)
        choice_from_id = _rating_choice_by_id(choices, selected_rating_id)
        if descriptor is not None and choice_from_descriptor is None:
            errors.append(
                f"{worksheet.title}!{row_number} has unknown Selected_Descriptor '{descriptor}' "
                f"for criterion '{criterion_id}'."
            )
        if selected_rating_id is not None and choice_from_id is None:
            errors.append(
                f"{worksheet.title}!{row_number} has unknown Selected_Rating_ID "
                f"'{selected_rating_id}' for criterion '{criterion_id}'."
            )
        if (
            choice_from_descriptor is not None
            and choice_from_id is not None
            and choice_from_descriptor.rating_id != choice_from_id.rating_id
        ):
            errors.append(
                f"{worksheet.title}!{row_number} Selected_Descriptor and Selected_Rating_ID "
                "refer to different ratings."
            )
        choice = choice_from_id or choice_from_descriptor
        if choice is None and points is None:
            errors.append(
                f"{worksheet.title}!{row_number} Action SET needs a selected descriptor, "
                "Selected_Rating_ID, or numeric Points."
            )
        if choice is not None and points is None:
            points = _parse_canvas_points(
                choice.points,
                location=f"Canvas rating '{choice.rating_id}'",
                errors=errors,
            )

        assessment: dict[str, Any] = {}
        if choice is not None:
            assessment["rating_id"] = choice.rating_id
        if points is not None:
            assessment["points"] = points
        if comment is not None:
            assessment["comments"] = comment
        inputs.append(
            _RubricInput(
                user_id=user_id,
                criterion_id=criterion_id,
                assessment=assessment or None,
                source_row=row_number,
            )
        )
    return inputs, errors


def _validate_input_students(
    *,
    grade_inputs: dict[int, _GradeInput],
    rubric_inputs: list[_RubricInput],
    submissions_by_user: dict[int, dict[str, Any]],
    errors: list[str],
) -> None:
    for user_id, grade_input in grade_inputs.items():
        if user_id not in submissions_by_user:
            errors.append(
                f"Grades row {grade_input.source_row} references Canvas user ID {user_id}, "
                "which is not a submission recipient for this assignment."
            )
    for rubric_input in rubric_inputs:
        if rubric_input.user_id not in grade_inputs:
            errors.append(
                f"Rubric row {rubric_input.source_row} references Canvas user ID "
                f"{rubric_input.user_id}, which is missing from the Grades sheet."
            )
        elif rubric_input.user_id not in submissions_by_user:
            errors.append(
                f"Rubric row {rubric_input.source_row} references Canvas user ID "
                f"{rubric_input.user_id}, which is not a submission recipient for this assignment."
            )


def _validate_snapshot_conflicts(
    *,
    grade_inputs: dict[int, _GradeInput],
    rubric_inputs: list[_RubricInput],
    snapshot: _GradebookSnapshot,
    submissions_by_user: dict[int, dict[str, Any]],
    errors: list[str],
) -> None:
    for grade_input in grade_inputs.values():
        submission = submissions_by_user.get(grade_input.user_id)
        if submission is None:
            continue
        if grade_input.grade is not None:
            expected_grade = snapshot.grades.get(grade_input.user_id)
            if grade_input.user_id not in snapshot.grades:
                errors.append(
                    f"Canvas Snapshot is missing the grade baseline for Canvas user ID "
                    f"{grade_input.user_id}."
                )
            elif not _same_canvas_value(expected_grade, _submission_grade(submission)):
                errors.append(
                    f"Grades row {grade_input.source_row} cannot be uploaded because the student's "
                    "grade changed in Canvas after this template was created."
                )
        if grade_input.submission_comment is not None:
            expected_submission = snapshot.submissions.get(grade_input.user_id)
            if expected_submission is None:
                errors.append(
                    f"Canvas Snapshot is missing the submission baseline for Canvas user ID "
                    f"{grade_input.user_id}."
                )
            elif not _same_submission_state(expected_submission, submission):
                errors.append(
                    f"Grades row {grade_input.source_row} cannot be uploaded because the student's "
                    "submission changed in Canvas after this template was created."
                )

    for rubric_input in rubric_inputs:
        if rubric_input.assessment is None:
            continue
        key = (rubric_input.user_id, rubric_input.criterion_id)
        expected = snapshot.rubric_assessments.get(key)
        if expected is None:
            errors.append(
                f"Canvas Snapshot is missing the rubric baseline for Canvas user ID "
                f"{rubric_input.user_id}, criterion '{rubric_input.criterion_id}'."
            )
            continue
        submission = submissions_by_user.get(rubric_input.user_id)
        if submission is None:
            continue
        expected_submission = snapshot.submissions.get(rubric_input.user_id)
        if expected_submission is not None and not _same_submission_state(
            expected_submission, submission
        ):
            errors.append(
                f"Rubric row {rubric_input.source_row} cannot be uploaded because the student's "
                "submission changed in Canvas after this template was created."
            )
            continue
        current = _submission_rubric_selection(submission, rubric_input.criterion_id)
        current_snapshot = _RubricSnapshot(
            rating_id=current.get("rating_id"),
            points=current.get("points"),
            comment=current.get("comments"),
        )
        if not (
            _same_canvas_value(expected.rating_id, current_snapshot.rating_id)
            and _same_canvas_value(expected.points, current_snapshot.points)
            and _same_canvas_value(expected.comment, current_snapshot.comment)
        ):
            errors.append(
                f"Rubric row {rubric_input.source_row} cannot be uploaded because the "
                "criterion assessment changed in Canvas after this template was created."
            )


def _read_mapped_values(
    worksheet: Any,
    row_number: int,
    columns: dict[str, int],
    fields: tuple[str, ...],
    errors: list[str],
) -> dict[str, object]:
    values: dict[str, object] = {}
    for field in fields:
        column = columns.get(field)
        if column is None:
            values[field] = None
            continue
        cell = worksheet.cell(row=row_number, column=column)
        if cell.data_type == "f":
            errors.append(
                f"{worksheet.title}!{cell.coordinate} ({field}) contains a formula. "
                "Enter a literal value before uploading."
            )
        values[field] = cell.value
    return values


def _parse_action(
    value: object,
    *,
    allowed: tuple[str, ...],
    location: str,
    errors: list[str],
) -> str | None:
    normalized = _text_or_none(value)
    if normalized is None:
        errors.append(f"{location} must be one of: {', '.join(allowed)}.")
        return None
    normalized = normalized.upper()
    if normalized not in allowed:
        errors.append(f"{location} must be one of: {', '.join(allowed)}.")
        return None
    return normalized


def _parse_grade_value(
    value: object,
    *,
    location: str,
    errors: list[str],
) -> str | int | float | None:
    if _is_blank(value):
        errors.append(f"{location} is required when Grade_Action is SET.")
        return None
    if isinstance(value, bool):
        errors.append(f"{location} must be a grade value, not a boolean.")
        return None
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not math.isfinite(value):
            errors.append(f"{location} must be finite.")
            return None
        return value
    if isinstance(value, str):
        return value.strip()
    errors.append(f"{location} must be text or a number.")
    return None


def _parse_optional_points(
    value: object,
    *,
    location: str,
    errors: list[str],
) -> int | float | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        errors.append(f"{location} must be numeric, not a boolean.")
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if math.isfinite(value):
            return value
        errors.append(f"{location} must be finite.")
        return None
    if isinstance(value, str):
        try:
            parsed = float(value.strip())
        except ValueError:
            errors.append(f"{location} must be numeric.")
            return None
        if not math.isfinite(parsed):
            errors.append(f"{location} must be finite.")
            return None
        return int(parsed) if parsed.is_integer() else parsed
    errors.append(f"{location} must be numeric.")
    return None


def _parse_canvas_points(
    value: object,
    *,
    location: str,
    errors: list[str],
) -> int | float | None:
    return _parse_optional_points(value, location=location, errors=errors)


def _submission_rubric_selection(
    submission: dict[str, Any],
    criterion_id: str,
) -> dict[str, Any]:
    rubric_assessment = submission.get("rubric_assessment")
    if not isinstance(rubric_assessment, dict):
        return {}
    for criterion_key in criterion_id_aliases(criterion_id):
        selection = rubric_assessment.get(criterion_key)
        if isinstance(selection, dict):
            return selection
    return {}


def _rating_choice_by_id(
    choices: list[_RatingChoice],
    rating_id: str | None,
) -> _RatingChoice | None:
    if rating_id is None:
        return None
    return next((choice for choice in choices if choice.rating_id == rating_id), None)


def _rating_choice_by_label(
    choices: list[_RatingChoice],
    selection_label: str | None,
) -> _RatingChoice | None:
    if selection_label is None:
        return None
    return next((choice for choice in choices if choice.selection_label == selection_label), None)


def _required_sheet(workbook: Workbook, sheet_name: str) -> Any:
    if sheet_name not in workbook.sheetnames:
        raise GradebookValidationError([f"Workbook is missing required sheet '{sheet_name}'."])
    return workbook[sheet_name]


def _header_index(worksheet: Any) -> dict[str, int]:
    headers: dict[str, int] = {}
    errors: list[str] = []
    for column_number in range(1, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=column_number).value
        header = _text_or_none(value)
        if header is None:
            continue
        normalized = header.casefold()
        if normalized in headers:
            errors.append(f"Sheet '{worksheet.title}' has duplicate column header '{header}'.")
            continue
        headers[normalized] = column_number
    if errors:
        raise GradebookValidationError(errors)
    return headers


def _required_header(headers: dict[str, int], header: str, sheet_name: str) -> int:
    column = headers.get(header.casefold())
    if column is None:
        raise GradebookValidationError([f"Sheet '{sheet_name}' is missing required column '{header}'."])
    return column


def _mapped_header_name(worksheet: Any, columns: dict[str, int], field: str) -> str:
    column = columns[field]
    return str(worksheet.cell(row=1, column=column).value)


def _submission_user(submission: dict[str, Any]) -> dict[str, Any]:
    user = submission.get("user")
    return user if isinstance(user, dict) else {}


def _submission_student_name(submission: dict[str, Any]) -> str | None:
    name = _submission_user(submission).get("name")
    return _text_or_none(name)


def _submission_grade(submission: dict[str, Any]) -> object:
    entered_grade = submission.get("entered_grade")
    return entered_grade if not _is_blank(entered_grade) else submission.get("grade")


def _assignment_name(assignment: dict[str, Any]) -> str:
    name = _text_or_none(assignment.get("name"))
    return name or "assignment"


def _student_label(user_id: int, student_name: str | None) -> str:
    return f"student {user_id} ({student_name})" if student_name else f"student {user_id}"


def _report(callback: Callable[[str], None] | None, message: str) -> None:
    if callback is not None:
        callback(message)


def _parse_positive_id(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float):
        if math.isfinite(value) and value.is_integer() and value > 0:
            return int(value)
        return None
    if isinstance(value, str) and re.fullmatch(r"[1-9]\d*", value.strip()):
        return int(value.strip())
    return None


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text_or_none(value: object) -> str | None:
    if _is_blank(value):
        return None
    return str(value).strip()


def _same_canvas_value(left: object, right: object) -> bool:
    return _normalise_compare_value(left) == _normalise_compare_value(right)


def _normalise_compare_value(value: object) -> tuple[str, str] | None:
    if _is_blank(value):
        return None
    if isinstance(value, bool):
        return ("text", str(value).casefold())
    if isinstance(value, (int, float, Decimal)):
        try:
            return ("number", str(Decimal(str(value)).normalize()))
        except (InvalidOperation, ValueError):
            return ("text", str(value).strip())
    if isinstance(value, str):
        text = value.strip()
        try:
            return ("number", str(Decimal(text).normalize()))
        except InvalidOperation:
            return ("text", text)
    return ("text", str(value).strip())


def _normalise_fingerprint_value(value: object) -> str | None:
    normalized = _normalise_compare_value(value)
    return normalized[1] if normalized is not None else None


def _format_points(value: object) -> str:
    normalized = _normalise_compare_value(value)
    if normalized is None:
        return "No default points"
    return normalized[1]


def _slugify_filename_part(value: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", value.strip()).strip("_")
    return slug.lower() or "assignment"
