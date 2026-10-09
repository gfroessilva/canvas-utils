"""Student roster export logic for Canvas courses."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from canvas_utils.api_client import CanvasClient

STUDENT_HEADERS = ["Student", "Email", "Student Number", "ID", "Groups"]


def export_course_students(
    client: CanvasClient,
    course_id: int,
    output: str | Path = ".",
    *,
    active_only: bool = False,
) -> list[Path]:
    """Export a course student roster to XLSX."""
    student_records = client.list_course_students(course_id, active_only=active_only)
    memberships = _build_group_memberships(client.list_course_groups(course_id))
    rows = [_map_student_row(record, memberships) for record in student_records]

    xlsx_path = _resolve_xlsx_output_path(Path(output), course_id)
    _write_workbook(xlsx_path, rows)
    return [xlsx_path]


def _map_student_row(student: dict[str, Any], memberships: dict[int, str]) -> dict[str, Any]:
    return {
        "Student": _string_or_none(student.get("name")),
        "Email": _string_or_none(student.get("email")) or _string_or_none(student.get("login_id")),
        "Student Number": _string_or_none(student.get("integration_id")),
        "ID": _coerce_positive_int(student.get("sis_user_id")),
        "Groups": memberships.get(_coerce_positive_int(student.get("id")) or -1),
    }


def _build_group_memberships(groups: list[dict[str, Any]]) -> dict[int, str]:
    memberships: dict[int, set[str]] = {}
    for group in groups:
        group_name = _string_or_none(group.get("name"))
        users = group.get("users")
        if group_name is None or not isinstance(users, list):
            continue

        for user in users:
            if not isinstance(user, dict):
                continue
            user_id = _coerce_positive_int(user.get("id"))
            if user_id is not None:
                memberships.setdefault(user_id, set()).add(group_name)

    return {
        user_id: "; ".join(sorted(group_names, key=str.casefold))
        for user_id, group_names in memberships.items()
    }


def _resolve_xlsx_output_path(output_path: Path, course_id: int) -> Path:
    export_date = datetime.now().strftime("%Y-%m-%d")
    file_name = f"canvas_students_{course_id}_{export_date}.xlsx"

    if output_path.suffix.lower() == ".xlsx":
        output_path.parent.mkdir(parents=True, exist_ok=True)
        return output_path

    output_path.mkdir(parents=True, exist_ok=True)
    return output_path / file_name


def _write_workbook(xlsx_path: Path, rows: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Students"

    sheet.append(STUDENT_HEADERS)
    for row in rows:
        sheet.append([row.get(header) for header in STUDENT_HEADERS])

    _autosize_columns(sheet, STUDENT_HEADERS)
    workbook.save(xlsx_path)


def _autosize_columns(sheet: Any, headers: list[str]) -> None:
    for index, header in enumerate(headers, start=1):
        max_length = len(header)
        for row_index in range(2, sheet.max_row + 1):
            cell_value = sheet.cell(row=row_index, column=index).value
            cell_length = len(str(cell_value)) if cell_value is not None else 0
            if cell_length > max_length:
                max_length = cell_length
        sheet.column_dimensions[get_column_letter(index)].width = max_length + 2


def _string_or_none(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _coerce_positive_int(value: object) -> int | None:
    if isinstance(value, int) and value > 0:
        return value
    if isinstance(value, str) and value.isdigit():
        parsed = int(value)
        return parsed if parsed > 0 else None
    return None
