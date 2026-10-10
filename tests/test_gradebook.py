from __future__ import annotations

from copy import deepcopy
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from canvas_utils.api_client import CanvasAPIError, CanvasTransportError
from canvas_utils.gradebook import (
    GRADE_HEADERS,
    GRADES_SHEET,
    MAPPING_SHEET,
    METADATA_SHEET,
    RUBRIC_CATALOG_SHEET,
    RUBRIC_HEADERS,
    RUBRIC_SHEET,
    SNAPSHOT_SHEET,
    GradebookUploadJournal,
    GradebookUploadLock,
    GradebookUploadLockError,
    GradebookValidationError,
    apply_gradebook_plan,
    build_gradebook_plan,
    create_gradebook_template,
    default_gradebook_upload_journal_path,
    default_gradebook_upload_lock_path,
)
from run import build_parser


class FakeGradebookClient:
    def __init__(self) -> None:
        self.assignment = {
            "id": 101,
            "name": "Assessment 1: Argument",
            "rubric": [
                {
                    "id": "_argument",
                    "description": "Argument quality",
                    "ratings": [
                        {"id": "rating-excellent", "description": "Excellent", "points": 5},
                        {"id": "rating-developing", "description": "Developing", "points": 3},
                    ],
                },
                {
                    "id": "_evidence",
                    "description": "Evidence",
                    "ratings": [
                        {"id": "rating-excellent-evidence", "description": "Excellent", "points": 5},
                    ],
                },
            ],
        }
        self.submissions = [
            {
                "assignment_id": 101,
                "user_id": 7,
                "attempt": 1,
                "submitted_at": "2026-09-10T09:00:00Z",
                "entered_grade": "70",
                "grade": "70",
                "user": {"name": "Jane Doe", "login_id": "jdoe"},
                "rubric_assessment": {
                    "criterion_argument": {
                        "rating_id": "rating-developing",
                        "points": 3,
                        "comments": "Needs a clearer thesis.",
                    }
                },
            },
            {
                "assignment_id": 101,
                "user_id": 8,
                "attempt": 1,
                "submitted_at": "2026-09-10T10:00:00Z",
                "entered_grade": None,
                "grade": None,
                "user": {"name": "John Roe", "login_id": "jroe"},
                "rubric_assessment": {},
            },
        ]
        self.assignment_requests: list[tuple[int, int, bool]] = []
        self.submission_requests: list[tuple[int, int, bool, bool]] = []
        self.single_submission_requests: list[tuple[int, int, int, bool]] = []
        self.grade_requests: list[dict[str, object]] = []
        self.failed_user_ids: set[int] = set()
        self.indeterminate_user_ids: set[int] = set()

    def get_assignment(
        self,
        course_id: int,
        assignment_id: int,
        *,
        include_rubric: bool = False,
    ) -> dict[str, object]:
        self.assignment_requests.append((course_id, assignment_id, include_rubric))
        return deepcopy(self.assignment)

    def list_submissions(
        self,
        course_id: int,
        assignment_id: int,
        include_user: bool = True,
        *,
        include_rubric_assessment: bool = False,
    ) -> list[dict[str, object]]:
        self.submission_requests.append(
            (course_id, assignment_id, include_user, include_rubric_assessment)
        )
        return deepcopy(self.submissions)

    def get_submission(
        self,
        course_id: int,
        assignment_id: int,
        user_id: int,
        *,
        include_rubric_assessment: bool = False,
    ) -> dict[str, object]:
        self.single_submission_requests.append(
            (course_id, assignment_id, user_id, include_rubric_assessment)
        )
        for submission in self.submissions:
            if submission["user_id"] == user_id:
                return deepcopy(submission)
        raise CanvasAPIError("Canvas submission not found.", status_code=404)

    def grade_submission(
        self,
        course_id: int,
        assignment_id: int,
        user_id: int,
        *,
        posted_grade: str | int | float | None = None,
        submission_comment: str | None = None,
        submission_comment_attempt: int | None = None,
        rubric_assessment: dict[str, dict[str, object]] | None = None,
    ) -> dict[str, object]:
        if user_id in self.indeterminate_user_ids:
            raise CanvasTransportError(
                "Connection closed before Canvas returned a response.",
                request_may_have_succeeded=True,
            )
        if user_id in self.failed_user_ids:
            raise CanvasAPIError("Canvas rejected this submission.", status_code=422)
        self.grade_requests.append(
            {
                "course_id": course_id,
                "assignment_id": assignment_id,
                "user_id": user_id,
                "posted_grade": posted_grade,
                "submission_comment": submission_comment,
                "submission_comment_attempt": submission_comment_attempt,
                "rubric_assessment": rubric_assessment,
            }
        )
        return {"id": user_id}


class GradebookTests(unittest.TestCase):
    def test_template_writes_controlled_sheets_metadata_and_descriptor_dropdowns(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(output_path, data_only=False)

        self.assertEqual(
            workbook.sheetnames,
            [
                "Instructions",
                GRADES_SHEET,
                RUBRIC_SHEET,
                RUBRIC_CATALOG_SHEET,
                MAPPING_SHEET,
                METADATA_SHEET,
                SNAPSHOT_SHEET,
                "Rubric Lists",
            ],
        )
        self.assertEqual(list(workbook[GRADES_SHEET].values)[0], tuple(GRADE_HEADERS))
        self.assertEqual(
            list(workbook[GRADES_SHEET].values)[1],
            (7, "Jane Doe", "jdoe", "70", None, "KEEP", None, "KEEP"),
        )
        self.assertEqual(list(workbook[RUBRIC_SHEET].values)[0], tuple(RUBRIC_HEADERS))
        self.assertEqual(
            list(workbook[RUBRIC_SHEET].values)[1][:8],
            (
                7,
                "Jane Doe",
                "_argument",
                "Argument quality",
                "3 - Developing",
                3,
                "Needs a clearer thesis.",
                None,
            ),
        )
        self.assertEqual(
            list(workbook[RUBRIC_CATALOG_SHEET].values)[1],
            ("_argument", "Argument quality", "rating-excellent", "5 - Excellent", "Excellent", 5),
        )
        self.assertEqual(workbook[METADATA_SHEET].sheet_state, "hidden")
        self.assertEqual(workbook[SNAPSHOT_SHEET].sheet_state, "hidden")
        self.assertEqual(len(workbook[RUBRIC_SHEET].data_validations.dataValidation), 2)
        self.assertEqual(client.assignment_requests, [(123, 101, True)])
        self.assertEqual(client.submission_requests, [(123, 101, True, True)])

    def test_plan_combines_grade_comment_and_rubric_descriptor_selection(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            grades = workbook[GRADES_SHEET]
            grades["E2"] = 85
            grades["F2"] = "SET"
            grades["G2"] = "Clear and well-supported work."
            grades["H2"] = "ADD"
            rubric = workbook[RUBRIC_SHEET]
            rubric["H2"] = "5 - Excellent"
            rubric["L2"] = "SET"
            workbook.save(gradebook_path)

            plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )

        self.assertEqual(plan.total_students, 2)
        self.assertEqual(plan.unchanged_students, 1)
        self.assertEqual(len(plan.changes), 1)
        change = plan.changes[0]
        self.assertEqual(change.user_id, 7)
        self.assertEqual(change.student_name, "Jane Doe")
        self.assertEqual(change.posted_grade, 85)
        self.assertEqual(change.submission_comment, "Clear and well-supported work.")
        self.assertEqual(change.submission_comment_attempt, 1)
        self.assertEqual(
            change.rubric_assessment,
            {"_argument": {"rating_id": "rating-excellent", "points": 5}},
        )

    def test_mapping_sheet_can_point_to_renamed_grade_column(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            grades = workbook[GRADES_SHEET]
            grades["E1"] = "Final Mark"
            grades["E2"] = "82"
            grades["F2"] = "SET"
            mapping = workbook[MAPPING_SHEET]
            for row_number in range(2, mapping.max_row + 1):
                if mapping.cell(row=row_number, column=1).value == "grades.new_grade":
                    mapping.cell(row=row_number, column=3).value = "Final Mark"
            workbook.save(gradebook_path)

            plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )

        self.assertEqual(len(plan.changes), 1)
        self.assertEqual(plan.changes[0].posted_grade, "82")

    def test_plan_rejects_a_grade_when_canvas_changed_since_template_creation(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["E2"] = 85
            workbook[GRADES_SHEET]["F2"] = "SET"
            workbook.save(gradebook_path)
            client.submissions[0]["entered_grade"] = "71"

            with self.assertRaisesRegex(GradebookValidationError, "grade changed in Canvas"):
                build_gradebook_plan(
                    client=client,
                    course_id=123,
                    assignment_id=101,
                    input_path=gradebook_path,
                )

    def test_upload_dry_run_does_not_mutate_and_execute_continues_after_failure(self) -> None:
        client = FakeGradebookClient()
        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            grades = workbook[GRADES_SHEET]
            grades["E2"] = 85
            grades["F2"] = "SET"
            grades["G3"] = "Good work."
            grades["H3"] = "ADD"
            workbook.save(gradebook_path)
            plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )

            dry_run = apply_gradebook_plan(client=client, plan=plan, execute=False)
            self.assertEqual(dry_run.planned, 2)
            self.assertEqual(dry_run.applied, 0)
            self.assertEqual(client.grade_requests, [])

            client.failed_user_ids.add(8)
            journal = GradebookUploadJournal.load(default_gradebook_upload_journal_path(gradebook_path))
            executed = apply_gradebook_plan(
                client=client,
                plan=plan,
                execute=True,
                upload_journal=journal,
            )

        self.assertEqual(executed.planned, 2)
        self.assertEqual(executed.applied, 1)
        self.assertEqual(executed.failed, 1)
        self.assertEqual(len(client.grade_requests), 1)
        self.assertEqual(client.grade_requests[0]["user_id"], 7)

    def test_execute_rechecks_canvas_before_each_grade_write(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["E2"] = 85
            workbook[GRADES_SHEET]["F2"] = "SET"
            workbook.save(gradebook_path)
            plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )
            client.submissions[0]["entered_grade"] = "71"
            journal = GradebookUploadJournal.load(default_gradebook_upload_journal_path(gradebook_path))

            result = apply_gradebook_plan(
                client=client,
                plan=plan,
                execute=True,
                upload_journal=journal,
            )

        self.assertEqual(result.applied, 0)
        self.assertEqual(result.conflicted, 1)
        self.assertEqual(result.failed, 0)
        self.assertEqual(client.grade_requests, [])
        self.assertEqual(client.single_submission_requests[-1], (123, 101, 7, False))

    def test_execute_rechecks_comment_only_submission_before_writing(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["G3"] = "Good work."
            workbook[GRADES_SHEET]["H3"] = "ADD"
            workbook.save(gradebook_path)
            plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )
            client.submissions[1]["attempt"] = 2
            client.submissions[1]["submitted_at"] = "2026-09-11T10:00:00Z"
            journal = GradebookUploadJournal.load(default_gradebook_upload_journal_path(gradebook_path))

            result = apply_gradebook_plan(
                client=client,
                plan=plan,
                execute=True,
                upload_journal=journal,
            )

        self.assertEqual(result.applied, 0)
        self.assertEqual(result.conflicted, 1)
        self.assertEqual(client.grade_requests, [])
        self.assertEqual(client.single_submission_requests[-1], (123, 101, 8, False))

    def test_execute_rechecks_rubric_only_submission_before_writing(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            rubric = workbook[RUBRIC_SHEET]
            rubric["H2"] = "5 - Excellent"
            rubric["L2"] = "SET"
            workbook.save(gradebook_path)
            plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )
            client.submissions[0]["attempt"] = 2
            client.submissions[0]["submitted_at"] = "2026-09-11T10:00:00Z"
            journal = GradebookUploadJournal.load(default_gradebook_upload_journal_path(gradebook_path))

            result = apply_gradebook_plan(
                client=client,
                plan=plan,
                execute=True,
                upload_journal=journal,
            )

        self.assertEqual(result.applied, 0)
        self.assertEqual(result.conflicted, 1)
        self.assertEqual(client.grade_requests, [])
        self.assertEqual(client.single_submission_requests[-1], (123, 101, 7, True))

    def test_indeterminate_comment_is_journaled_and_not_replayed_by_default(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["G3"] = "Good work."
            workbook[GRADES_SHEET]["H3"] = "ADD"
            workbook.save(gradebook_path)
            plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )
            journal_path = default_gradebook_upload_journal_path(gradebook_path)
            journal = GradebookUploadJournal.load(journal_path)
            client.indeterminate_user_ids.add(8)

            first_result = apply_gradebook_plan(
                client=client,
                plan=plan,
                execute=True,
                upload_journal=journal,
            )
            reloaded_journal = GradebookUploadJournal.load(journal_path)
            second_result = apply_gradebook_plan(
                client=client,
                plan=plan,
                execute=True,
                upload_journal=reloaded_journal,
            )

        self.assertEqual(first_result.indeterminate, 1)
        self.assertEqual(first_result.failed, 0)
        self.assertEqual(second_result.applied, 0)
        self.assertEqual(second_result.journal_skipped, 1)
        self.assertEqual(client.grade_requests, [])

    def test_pending_journal_record_survives_a_post_write_journal_failure(self) -> None:
        client = FakeGradebookClient()

        class FailingCompletionJournal(GradebookUploadJournal):
            def complete(self, outcome: object) -> None:
                raise OSError("disk full")

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["E2"] = 85
            workbook[GRADES_SHEET]["F2"] = "SET"
            workbook.save(gradebook_path)
            plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )
            journal_path = default_gradebook_upload_journal_path(gradebook_path)
            journal = FailingCompletionJournal(journal_path, {})

            result = apply_gradebook_plan(
                client=client,
                plan=plan,
                execute=True,
                upload_journal=journal,
            )
            persisted_journal = GradebookUploadJournal.load(journal_path)

        self.assertEqual(result.applied, 1)
        self.assertEqual(result.journal_errors, 1)
        self.assertEqual(len(client.grade_requests), 1)
        self.assertEqual(len(persisted_journal.excluded_change_ids()), 1)

    def test_new_template_can_repeat_a_prior_payload_with_the_same_journal_path(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = Path(temp_dir) / "gradebook.xlsx"
            create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=gradebook_path,
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["E2"] = 80
            workbook[GRADES_SHEET]["F2"] = "SET"
            workbook.save(gradebook_path)
            first_plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )
            journal = GradebookUploadJournal.load(default_gradebook_upload_journal_path(gradebook_path))
            first_result = apply_gradebook_plan(
                client=client,
                plan=first_plan,
                execute=True,
                upload_journal=journal,
            )
            client.submissions[0]["entered_grade"] = "80"
            client.submissions[0]["grade"] = "80"

            client.submissions[0]["entered_grade"] = "70"
            client.submissions[0]["grade"] = "70"
            create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=gradebook_path,
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["E2"] = 80
            workbook[GRADES_SHEET]["F2"] = "SET"
            workbook.save(gradebook_path)
            second_plan = build_gradebook_plan(
                client=client,
                course_id=123,
                assignment_id=101,
                input_path=gradebook_path,
            )
            second_result = apply_gradebook_plan(
                client=client,
                plan=second_plan,
                execute=True,
                upload_journal=journal,
            )

        self.assertEqual(first_result.applied, 1)
        self.assertEqual(second_result.applied, 1)
        self.assertEqual(second_result.journal_skipped, 0)
        self.assertEqual(len(client.grade_requests), 2)

    def test_rejects_malformed_workbook_and_prevents_parallel_uploads(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            malformed_path = Path(temp_dir) / "not-a-workbook.xlsx"
            malformed_path.write_text("not an XLSX workbook", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "Could not read gradebook workbook"):
                build_gradebook_plan(
                    client=client,
                    course_id=123,
                    assignment_id=101,
                    input_path=malformed_path,
                )

            journal_path = Path(temp_dir) / "gradebook.upload-journal.json"
            lock_path = default_gradebook_upload_lock_path(journal_path)
            with GradebookUploadLock(lock_path):
                with self.assertRaises(GradebookUploadLockError):
                    with GradebookUploadLock(lock_path):
                        pass

    def test_plan_rejects_grade_values_incompatible_with_assignment_grading_type(self) -> None:
        client = FakeGradebookClient()
        client.assignment["grading_type"] = "pass_fail"
        client.assignment["points_possible"] = 10

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["E2"] = 5
            workbook[GRADES_SHEET]["F2"] = "SET"
            workbook.save(gradebook_path)

            with self.assertRaisesRegex(GradebookValidationError, "pass/fail assignment"):
                build_gradebook_plan(
                    client=client,
                    course_id=123,
                    assignment_id=101,
                    input_path=gradebook_path,
                )

        client = FakeGradebookClient()
        client.assignment["grading_type"] = "not_graded"
        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            workbook[GRADES_SHEET]["E2"] = 85
            workbook[GRADES_SHEET]["F2"] = "SET"
            workbook.save(gradebook_path)

            with self.assertRaisesRegex(GradebookValidationError, "not graded"):
                build_gradebook_plan(
                    client=client,
                    course_id=123,
                    assignment_id=101,
                    input_path=gradebook_path,
                )

    def test_plan_rejects_unknown_rubric_descriptor_before_upload(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            rubric = workbook[RUBRIC_SHEET]
            rubric["H2"] = "Not a Canvas rating"
            rubric["L2"] = "SET"
            workbook.save(gradebook_path)

            with self.assertRaisesRegex(GradebookValidationError, "unknown Selected_Descriptor"):
                build_gradebook_plan(
                    client=client,
                    course_id=123,
                    assignment_id=101,
                    input_path=gradebook_path,
                )

    def test_cli_exposes_gradebook_actions_and_explicit_execute(self) -> None:
        parser = build_parser()

        template_args = parser.parse_args(
            ["gradebook", "template", "--course", "123", "--assignment", "101"]
        )
        upload_args = parser.parse_args(
            [
                "gradebook",
                "upload",
                "--course",
                "123",
                "--assignment",
                "101",
                "--input",
                "gradebook.xlsx",
                "--execute",
                "--journal",
                "journal.json",
                "--retry-indeterminate",
            ]
        )

        self.assertEqual(template_args.command, "gradebook")
        self.assertEqual(template_args.gradebook_action, "template")
        self.assertEqual(template_args.course, 123)
        self.assertEqual(template_args.assignment, 101)
        self.assertEqual(upload_args.gradebook_action, "upload")
        self.assertTrue(upload_args.execute)
        self.assertEqual(upload_args.journal, "journal.json")
        self.assertTrue(upload_args.retry_indeterminate)

    def test_mapping_rejects_duplicate_target_columns(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            mapping_sheet = workbook["Canvas Mapping"]
            # Row 5 is grades.submission_comment -> New_Submission_Comment
            # Change its target column to New_Grade (which row 3 grades.new_grade already targets)
            mapping_sheet["C5"] = "New_Grade"
            workbook.save(gradebook_path)

            with self.assertRaisesRegex(GradebookValidationError, "maps multiple fields .* to the same column"):
                build_gradebook_plan(
                    client=client,
                    course_id=123,
                    assignment_id=101,
                    input_path=gradebook_path,
                )

    def test_create_gradebook_template_sanitizes_formula_injection(self) -> None:
        class FakeFormulaGradebookClient(FakeGradebookClient):
            def list_submissions(
                self,
                course_id: int,
                assignment_id: int,
                include_user: bool = True,
                *,
                include_rubric_assessment: bool = False,
            ) -> list[dict[str, object]]:
                return [
                    {
                        "assignment_id": assignment_id,
                        "user_id": 7,
                        "score": 10,
                        "grade": "=1+1",
                        "workflow_state": "graded",
                        "user": {"name": "=SUM(A1:A5)", "login_id": "=cmd|' /C calc'!A0"},
                        "rubric_assessment": {},
                    }
                ]

        client = FakeFormulaGradebookClient()
        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path, data_only=False)
            grades = workbook[GRADES_SHEET]
            self.assertEqual(grades["B2"].value, "'=SUM(A1:A5)")
            self.assertEqual(grades["C2"].value, "'=cmd|' /C calc'!A0")
            self.assertEqual(grades["D2"].value, "'=1+1")

    def test_plan_rejects_rubric_upload_when_submission_attempt_changes(self) -> None:
        client = FakeGradebookClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            gradebook_path = create_gradebook_template(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )
            workbook = load_workbook(gradebook_path)
            rubric = workbook[RUBRIC_SHEET]
            rubric["H2"] = "5 - Excellent"
            rubric["L2"] = "SET"
            workbook.save(gradebook_path)

            # Change submission attempt in client
            client.submissions[0]["attempt"] = 2

            with self.assertRaisesRegex(GradebookValidationError, "submission changed in Canvas after this template was created"):
                build_gradebook_plan(
                    client=client,
                    course_id=123,
                    assignment_id=101,
                    input_path=gradebook_path,
                )


if __name__ == "__main__":
    unittest.main()
