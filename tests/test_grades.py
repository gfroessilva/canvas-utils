from __future__ import annotations

import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from canvas_utils.api_client import CanvasAPIError
from canvas_utils.grades import GRADE_HEADERS, MARKER_STATS_HEADERS, export_assignment_grades
from run import build_parser


class FakeGradesClient:
    def __init__(self) -> None:
        self.user_requests: list[int] = []
        self.course_user_requests: list[tuple[int, int]] = []

    def list_assignments(self, course_id: int) -> list[dict[str, object]]:
        return [{"id": 101, "name": "Assessment 1: Intro / Basics"}]

    def list_submissions(
        self,
        course_id: int,
        assignment_id: int,
        include_user: bool = True,
    ) -> list[dict[str, object]]:
        return [
            {
                "assignment_id": assignment_id,
                "user_id": 7,
                "score": 88.5,
                "grade": "88.5",
                "entered_grade": "88.5",
                "workflow_state": "graded",
                "grader_id": 99,
                "user": {"name": "Jane Doe", "login_id": "jdoe"},
            },
            {
                "assignment_id": assignment_id,
                "user_id": 8,
                "score": 79,
                "grade": "79",
                "entered_grade": "79",
                "workflow_state": "graded",
                "grader_id": "99",
                "user": {"name": "John Roe", "login_id": "jroe"},
            },
            {
                "assignment_id": assignment_id,
                "user_id": 9,
                "score": 10,
                "grade": "complete",
                "entered_grade": "complete",
                "workflow_state": "graded",
                "grader_id": -42,
                "user": {"name": "Auto Grade", "login_id": "agrade"},
            },
            {
                "assignment_id": assignment_id,
                "user_id": 10,
                "score": None,
                "grade": None,
                "entered_grade": None,
                "workflow_state": "submitted",
                "grader_id": None,
                "user": {"name": "Waiting Grade", "login_id": "wgrade"},
            },
        ]

    def get_user(self, user_id: int) -> dict[str, object]:
        self.user_requests.append(user_id)
        return {
            "id": user_id,
            "name": "Marker One",
            "login_id": "marker1",
            "sis_user_id": "M001",
        }

    def get_course_user(self, course_id: int, user_id: int) -> dict[str, object]:
        self.course_user_requests.append((course_id, user_id))
        return {
            "id": user_id,
            "name": "Course Marker",
            "login_id": "coursemarker",
            "sis_user_id": "CM001",
        }


class FakeCourseScopedGraderClient(FakeGradesClient):
    def list_submissions(
        self,
        course_id: int,
        assignment_id: int,
        include_user: bool = True,
    ) -> list[dict[str, object]]:
        return [
            {
                "assignment_id": assignment_id,
                "user_id": 7,
                "score": 88.5,
                "grade": "88.5",
                "entered_grade": "88.5",
                "workflow_state": "graded",
                "grader_id": 74955,
                "user": {"name": "Jane Doe", "login_id": "jdoe"},
            }
        ]

    def get_user(self, user_id: int) -> dict[str, object]:
        self.user_requests.append(user_id)
        raise CanvasAPIError(
            "Canvas resource not found (404 Not Found).",
            status_code=404,
        )


class FakeUnresolvableGraderClient(FakeCourseScopedGraderClient):
    def get_course_user(self, course_id: int, user_id: int) -> dict[str, object]:
        self.course_user_requests.append((course_id, user_id))
        raise CanvasAPIError(
            "Canvas resource not found (404 Not Found).",
            status_code=404,
        )


class GradeExportTests(unittest.TestCase):
    def test_export_assignment_grades_writes_xlsx_with_grades_and_marker_stats(self) -> None:
        client = FakeGradesClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            written_files = export_assignment_grades(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )

            self.assertEqual(len(written_files), 1)
            xlsx_path = written_files[0]
            self.assertEqual(
                bool(re.fullmatch(r"canvas_grades_101_assessment_1_intro_basics_\d{4}-\d{2}-\d{2}\.xlsx", xlsx_path.name)),
                True,
            )

            workbook = load_workbook(xlsx_path, data_only=False)
            grades_sheet = workbook["Grades"]
            stats_sheet = workbook["Marker Statistics"]
            grade_rows = list(grades_sheet.iter_rows(values_only=True))
            stats_rows = list(stats_sheet.iter_rows(values_only=True))

        self.assertEqual(list(grade_rows[0]), GRADE_HEADERS)
        self.assertEqual(client.user_requests, [99])
        self.assertEqual(
            grade_rows[1],
            (
                101,
                "Assessment 1: Intro / Basics",
                7,
                "Jane Doe",
                "jdoe",
                88.5,
                "88.5",
                "88.5",
                "graded",
                99,
                "Marker One",
                "marker1",
                "M001",
                "human",
            ),
        )
        self.assertEqual(grade_rows[3][10], "Auto-graded")
        self.assertEqual(grade_rows[3][13], "automatic")
        self.assertEqual(grade_rows[4][10], "Ungraded")
        self.assertEqual(grade_rows[4][13], "ungraded")
        self.assertEqual(grade_rows[4][8], "submitted")

        self.assertEqual(list(stats_rows[0]), MARKER_STATS_HEADERS)
        stats_by_name = {row[1]: row for row in stats_rows[1:]}
        stats_row_numbers = {row[1]: index for index, row in enumerate(stats_rows[1:], start=2)}
        marker_one_row = stats_row_numbers["Marker One"]
        auto_graded_row = stats_row_numbers["Auto-graded"]
        ungraded_row = stats_row_numbers["Ungraded"]
        marker_one_stats = stats_by_name["Marker One"]
        self.assertEqual(marker_one_stats[:5], (99, "Marker One", "marker1", "M001", "human"))
        self.assertEqual(
            marker_one_stats[5],
            f'=COUNTIFS(Grades!$J$2:$J$1048576,A{marker_one_row},Grades!$K$2:$K$1048576,B{marker_one_row},Grades!$L$2:$L$1048576,C{marker_one_row},Grades!$M$2:$M$1048576,D{marker_one_row},Grades!$N$2:$N$1048576,E{marker_one_row})',
        )
        self.assertIn(f'A{marker_one_row}', marker_one_stats[6])
        self.assertIn(f'B{marker_one_row}', marker_one_stats[6])
        self.assertIn("ISNUMBER(Grades!$F$2:$F$1048576)", marker_one_stats[6])
        self.assertEqual(marker_one_stats[7], f'=F{marker_one_row}-G{marker_one_row}')
        self.assertTrue(marker_one_stats[8].startswith(f'=IF(G{marker_one_row}=0,"",ROUND(AVERAGEIFS('))
        self.assertTrue(marker_one_stats[9].startswith(f'=IF(G{marker_one_row}=0,"",MINIFS('))
        self.assertTrue(marker_one_stats[10].startswith(f'=IF(G{marker_one_row}=0,"",MAXIFS('))

        auto_graded_stats = stats_by_name["Auto-graded"]
        self.assertEqual(auto_graded_stats[:5], (-42, "Auto-graded", None, None, "automatic"))
        self.assertEqual(
            auto_graded_stats[5],
            f'=COUNTIFS(Grades!$J$2:$J$1048576,A{auto_graded_row},Grades!$K$2:$K$1048576,B{auto_graded_row},Grades!$L$2:$L$1048576,C{auto_graded_row},Grades!$M$2:$M$1048576,D{auto_graded_row},Grades!$N$2:$N$1048576,E{auto_graded_row})',
        )
        self.assertIn(f'A{auto_graded_row}', auto_graded_stats[6])
        self.assertEqual(auto_graded_stats[7], f'=F{auto_graded_row}-G{auto_graded_row}')
        self.assertTrue(stats_by_name["Ungraded"][8].startswith(f'=IF(G{ungraded_row}=0,"",ROUND(AVERAGEIFS('))
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        self.assertTrue(workbook.calculation.forceFullCalc)

    def test_export_assignment_grades_falls_back_to_course_scoped_user_lookup(self) -> None:
        client = FakeCourseScopedGraderClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            written_files = export_assignment_grades(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )

            workbook = load_workbook(written_files[0], data_only=False)
            grade_rows = list(workbook["Grades"].iter_rows(values_only=True))

        self.assertEqual(client.user_requests, [74955])
        self.assertEqual(client.course_user_requests, [(123, 74955)])
        self.assertEqual(grade_rows[1][9], 74955)
        self.assertEqual(grade_rows[1][10], "Course Marker")
        self.assertEqual(grade_rows[1][11], "coursemarker")
        self.assertEqual(grade_rows[1][12], "CM001")
        self.assertEqual(grade_rows[1][13], "human")

    def test_export_assignment_grades_keeps_rows_when_grader_lookup_404s(self) -> None:
        client = FakeUnresolvableGraderClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            written_files = export_assignment_grades(
                client=client,
                course_id=123,
                assignment_id=101,
                output=Path(temp_dir),
            )

            workbook = load_workbook(written_files[0], data_only=False)
            grade_rows = list(workbook["Grades"].iter_rows(values_only=True))

        self.assertEqual(client.user_requests, [74955])
        self.assertEqual(client.course_user_requests, [(123, 74955)])
        self.assertEqual(grade_rows[1][9], 74955)
        self.assertEqual(grade_rows[1][10], "Unresolvable grader")
        self.assertEqual(grade_rows[1][13], "unresolvable")

    def test_cli_exposes_grades_subcommand(self) -> None:
        parser = build_parser()
        args = parser.parse_args(["grades", "--course", "123", "--assignment", "456"])
        self.assertEqual(args.command, "grades")
        self.assertEqual(args.course, 123)
        self.assertEqual(args.assignment, "456")


if __name__ == "__main__":
    unittest.main()
