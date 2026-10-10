"""Unit and integration tests for Canvas quiz authoring and deployment."""

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock

from canvas_utils.api_client import CanvasClient
from canvas_utils.quizzes import (
    AnswerSpec,
    QuestionSpec,
    QuizSpec,
    QuizValidationError,
    create_quiz_template,
    create_quiz_template_excel,
    create_quiz_template_json,
    deploy_quiz,
    load_quiz_from_dict,
    load_quiz_from_excel,
    load_quiz_from_json,
)


class TestQuizModelsAndValidation(unittest.TestCase):
    def test_valid_quiz_validation(self):
        quiz = QuizSpec(
            title="Intro Quiz",
            quiz_type="practice_quiz",
            questions=[
                QuestionSpec(
                    name="Q1",
                    text="Is this valid?",
                    question_type="true_false",
                    answers=[
                        AnswerSpec(text="True", weight=100.0),
                        AnswerSpec(text="False", weight=0.0),
                    ],
                )
            ],
        )
        quiz.validate()
        self.assertEqual(quiz.total_points, 1.0)

    def test_quiz_missing_title_raises(self):
        quiz = QuizSpec(
            title="  ",
            questions=[
                QuestionSpec(
                    name="Q1",
                    text="Prompt",
                    question_type="essay",
                )
            ],
        )
        with self.assertRaises(QuizValidationError) as ctx:
            quiz.validate()
        self.assertIn("Quiz title cannot be blank", str(ctx.exception))

    def test_quiz_no_questions_raises(self):
        quiz = QuizSpec(title="Empty Quiz", questions=[])
        with self.assertRaises(QuizValidationError):
            quiz.validate()

    def test_multiple_choice_requires_exact_single_correct(self):
        # 0 correct
        q_zero = QuestionSpec(
            name="Q1",
            text="MC Prompt",
            question_type="multiple_choice",
            answers=[
                AnswerSpec(text="A", weight=0.0),
                AnswerSpec(text="B", weight=0.0),
            ],
        )
        with self.assertRaises(QuizValidationError):
            q_zero.validate()

        # 2 correct
        q_two = QuestionSpec(
            name="Q1",
            text="MC Prompt",
            question_type="multiple_choice",
            answers=[
                AnswerSpec(text="A", weight=100.0),
                AnswerSpec(text="B", weight=100.0),
            ],
        )
        with self.assertRaises(QuizValidationError):
            q_two.validate()

    def test_multiple_answers_validation(self):
        q = QuestionSpec(
            name="Q1",
            text="Choose all that apply",
            question_type="multiple_answers",
            answers=[
                AnswerSpec(text="A", weight=50.0),
                AnswerSpec(text="B", weight=50.0),
                AnswerSpec(text="C", weight=0.0),
            ],
        )
        q.validate()
        self.assertEqual(len(q.answers), 3)


class TestQuizLoaders(unittest.TestCase):
    def test_load_from_json(self):
        payload = {
            "title": "JSON Quiz",
            "quiz_type": "practice_quiz",
            "time_limit": 30,
            "questions": [
                {
                    "name": "Q1",
                    "text": "What is 2+2?",
                    "type": "multiple_choice",
                    "points_possible": 2.0,
                    "answers": [
                        {"text": "3", "weight": 0},
                        {"text": "4", "weight": 100},
                    ],
                }
            ],
        }
        with tempfile.NamedTemporaryFile("w", suffix=".json", delete=False, encoding="utf-8") as f:
            json.dump(payload, f)
            temp_path = Path(f.name)

        try:
            quiz = load_quiz_from_json(temp_path)
            self.assertEqual(quiz.title, "JSON Quiz")
            self.assertEqual(quiz.quiz_type, "practice_quiz")
            self.assertEqual(quiz.time_limit, 30)
            self.assertEqual(len(quiz.questions), 1)
            self.assertEqual(quiz.questions[0].points_possible, 2.0)
        finally:
            temp_path.unlink(missing_ok=True)

    def test_create_and_load_excel_template(self):
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            temp_path = Path(f.name)

        try:
            create_quiz_template_excel(temp_path)
            quiz = load_quiz_from_excel(temp_path)
            self.assertEqual(quiz.title, "Sample Module Quiz")
            self.assertEqual(len(quiz.questions), 5)
            self.assertEqual(quiz.total_points, 10.0)

            # Check question types parsed correctly
            q_types = [q.question_type for q in quiz.questions]
            self.assertIn("multiple_choice_question", q_types)
            self.assertIn("true_false_question", q_types)
            self.assertIn("multiple_answers_question", q_types)
            self.assertIn("short_answer_question", q_types)
            self.assertIn("essay_question", q_types)
        finally:
            temp_path.unlink(missing_ok=True)

    def test_create_and_load_json_template(self):
        with tempfile.NamedTemporaryFile(suffix=".json", delete=False) as f:
            temp_path = Path(f.name)

        try:
            create_quiz_template_json(temp_path)
            quiz = load_quiz_from_json(temp_path)
            self.assertEqual(quiz.title, "Sample Module Quiz")
            self.assertEqual(len(quiz.questions), 5)
            self.assertEqual(quiz.total_points, 10.0)

            # Check question types parsed correctly
            q_types = [q.question_type for q in quiz.questions]
            self.assertIn("multiple_choice_question", q_types)
            self.assertIn("true_false_question", q_types)
            self.assertIn("multiple_answers_question", q_types)
            self.assertIn("short_answer_question", q_types)
            self.assertIn("essay_question", q_types)
        finally:
            temp_path.unlink(missing_ok=True)

    def test_create_quiz_template_auto_dispatch(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            json_target = Path(tmpdir) / "quiz.json"
            xlsx_target = Path(tmpdir) / "quiz.xlsx"
            create_quiz_template(json_target)
            create_quiz_template(xlsx_target)
            self.assertTrue(json_target.exists())
            self.assertTrue(xlsx_target.exists())


class TestQuizDeployment(unittest.TestCase):
    def test_deploy_quiz_creates_draft_and_publishes_at_end(self):
        mock_client = MagicMock(spec=CanvasClient)
        mock_client.create_quiz.return_value = {"id": 999, "html_url": "https://canvas/quizzes/999"}
        mock_client.create_quiz_question.return_value = {"id": 1001}
        mock_client.update_quiz.return_value = {"id": 999, "published": True}

        quiz_spec = QuizSpec(
            title="Deploy Test",
            questions=[
                QuestionSpec(
                    name="Q1",
                    text="T/F Question",
                    question_type="true_false",
                    general_feedback="General feedback note",
                    correct_comments="Good job",
                    incorrect_comments="Try again",
                    answers=[
                        AnswerSpec(text="True", weight=100.0, comments="Correct explanation"),
                        AnswerSpec(text="False", weight=0.0, comments="Incorrect explanation"),
                    ],
                )
            ],
        )

        result = deploy_quiz(mock_client, course_id=42, quiz_spec=quiz_spec, publish=True)

        self.assertEqual(result["quiz_id"], 999)
        self.assertEqual(result["question_count"], 1)

        # 1. Initial creation must be draft (published=False)
        mock_client.create_quiz.assert_called_once()
        create_quiz_args = mock_client.create_quiz.call_args[0]
        self.assertEqual(create_quiz_args[0], 42)
        self.assertFalse(create_quiz_args[1]["published"])

        # 2. Check question feedback payload includes Canvas expected keys
        mock_client.create_quiz_question.assert_called_once()
        q_args = mock_client.create_quiz_question.call_args[0]
        q_payload = q_args[2]
        self.assertEqual(q_payload["neutral_comments"], "General feedback note")
        self.assertEqual(q_payload["neutral_comments_html"], "General feedback note")
        self.assertEqual(q_payload["correct_comments_html"], "Good job")
        self.assertEqual(q_payload["incorrect_comments_html"], "Try again")
        self.assertEqual(q_payload["answers"][0]["answer_comments"], "Correct explanation")
        self.assertEqual(q_payload["answers"][0]["answer_comment_html"], "Correct explanation")

        # 3. Final publish update
        mock_client.update_quiz.assert_called_once_with(42, 999, {"published": True})


class TestCopilotEdgeCases(unittest.TestCase):
    def test_excel_zero_points_preserved(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws_q = wb.active
        ws_q.title = "Questions"
        ws_q.append(["Question_Name", "Question_Type", "Points", "Question_Text", "Choice_A", "Choice_B", "Correct_Answer"])
        ws_q.append(["Q0", "multiple_choice", 0, "Zero point question", "Choice 1", "Choice 2", "A"])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            temp_path = Path(f.name)

        try:
            wb.save(temp_path)
            quiz = load_quiz_from_excel(temp_path)
            self.assertEqual(quiz.questions[0].points_possible, 0.0)
            self.assertEqual(quiz.total_points, 0.0)
        finally:
            temp_path.unlink(missing_ok=True)

    def test_excel_invalid_points_raises(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws_q = wb.active
        ws_q.title = "Questions"
        ws_q.append(["Question_Name", "Question_Type", "Points", "Question_Text", "Choice_A", "Choice_B", "Correct_Answer"])
        ws_q.append(["Q1", "multiple_choice", "not_a_number", "Question", "Choice 1", "Choice 2", "A"])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            temp_path = Path(f.name)

        try:
            wb.save(temp_path)
            with self.assertRaises(QuizValidationError):
                load_quiz_from_excel(temp_path)
        finally:
            temp_path.unlink(missing_ok=True)

    def test_excel_unsupported_question_type_raises(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws_q = wb.active
        ws_q.title = "Questions"
        ws_q.append(["Question_Name", "Question_Type", "Points", "Question_Text", "Choice_A", "Choice_B", "Correct_Answer"])
        ws_q.append(["Q1", "unsupported_magic_type", 1.0, "Question", "Choice 1", "Choice 2", "A"])

        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            temp_path = Path(f.name)

        try:
            wb.save(temp_path)
            with self.assertRaises(QuizValidationError) as ctx:
                load_quiz_from_excel(temp_path)
            self.assertIn("Unsupported question type", str(ctx.exception))
        finally:
            temp_path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()
