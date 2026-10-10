from __future__ import annotations

import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from canvas_utils.students import STUDENT_HEADERS, export_course_students
from run import build_parser


class FakeStudentsClient:
    def __init__(self) -> None:
        self.calls: list[tuple[int, bool]] = []
        self.group_calls: list[int] = []

    def list_course_students(self, course_id: int, *, active_only: bool = False) -> list[dict[str, object]]:
        self.calls.append((course_id, active_only))
        return [
            {
                "id": 7,
                "name": "Jane Doe",
                "email": "jane@example.edu",
                "integration_id": "S12345",
                "sis_user_id": "12345",
                "login_id": "jdoe",
            },
            {
                "id": "8",
                "name": "John Roe",
                "email": None,
                "integration_id": "S98765",
                "sis_user_id": "98765",
                "login_id": "jroe@example.edu",
            },
        ]

    def list_course_groups(self, course_id: int) -> list[dict[str, object]]:
        self.group_calls.append(course_id)
        return [
            {"name": "Project Group B", "users": [{"id": "8"}]},
            {"name": "Project Group A", "users": [{"id": 7}, {"id": 8}]},
            {"name": "Empty Group", "users": []},
        ]


class StudentExportTests(unittest.TestCase):
    def test_export_course_students_writes_expected_xlsx_schema(self) -> None:
        client = FakeStudentsClient()
        with tempfile.TemporaryDirectory() as temp_dir:
            written_files = export_course_students(
                client=client,
                course_id=123,
                output=Path(temp_dir),
            )

            self.assertEqual(len(written_files), 1)
            xlsx_path = written_files[0]
            self.assertEqual(
                bool(re.fullmatch(r"canvas_students_123_\d{4}-\d{2}-\d{2}\.xlsx", xlsx_path.name)),
                True,
            )

            workbook = load_workbook(xlsx_path)
            sheet = workbook["Students"]
            rows = list(sheet.iter_rows(values_only=True))

        self.assertEqual(client.calls, [(123, False)])
        self.assertEqual(client.group_calls, [123])
        self.assertEqual(list(rows[0]), STUDENT_HEADERS)
        self.assertEqual(rows[1], ("Jane Doe", "jane@example.edu", "12345", 7, "Project Group A"))
        self.assertEqual(
            rows[2],
            ("John Roe", "jroe@example.edu", "98765", 8, "Project Group A; Project Group B"),
        )

    def test_export_course_students_accepts_explicit_xlsx_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "roster.xlsx"
            written_files = export_course_students(
                client=FakeStudentsClient(),
                course_id=987,
                output=output_path,
            )

            self.assertEqual(written_files, [output_path])
            self.assertTrue(output_path.exists())

    def test_cli_exposes_students_subcommand(self) -> None:
        parser = build_parser()
        args = parser.parse_args(["students", "--course", "123", "--output", "./exports", "--active-only"])

        self.assertEqual(args.command, "students")
        self.assertEqual(args.course, 123)
        self.assertEqual(args.output, "./exports")
        self.assertTrue(args.active_only)

    def test_export_course_students_forwards_active_only(self) -> None:
        client = FakeStudentsClient()

        with tempfile.TemporaryDirectory() as temp_dir:
            export_course_students(
                client=client,
                course_id=555,
                output=Path(temp_dir),
                active_only=True,
            )

        self.assertEqual(client.calls, [(555, True)])
        self.assertEqual(client.group_calls, [555])


if __name__ == "__main__":
    unittest.main()
